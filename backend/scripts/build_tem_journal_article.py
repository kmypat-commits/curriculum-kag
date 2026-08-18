"""Build English and Russian TEM Journal manuscripts from the official DOCX template."""

from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


METRIC_KEYS = {
    "provenance": "epvo_provenance",
    "profile_provenance": "epvo_provenance_excluding_regulatory",
    "raw_semester": "semester_alignment_pm1",
    "prereq_semester": "semester_alignment_prereq_adjusted_pm1",
    "semantic_semester": "semester_alignment_semantic_adjusted_pm1",
    "quality": "international_score",
    "hard": "hard_violations",
}


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def remove_after_front_matter(document: Document) -> None:
    body = document._body._element
    paragraphs = document.paragraphs
    boundary = paragraphs[9]._p
    seen = False
    for child in list(body):
        if child is boundary:
            seen = True
            continue
        if seen and child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, size=11, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


INLINE_TOKEN = re.compile(r"https?://[^\s]+|(?<![\w])([A-Za-z])_([A-Za-z0-9]{1,16})(?![\w])")


def add_hyperlink(paragraph, text: str, url: str, size=11):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), str(size * 2))
    for node in (fonts, color, underline, size_node):
        r_pr.append(node)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_formatted_text(paragraph, text: str, *, size=11, bold=False, italic=False):
    cursor = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            set_run_font(paragraph.add_run(text[cursor:match.start()]), size=size, bold=bold, italic=italic)
        token = match.group(0)
        if token.startswith(("http://", "https://")):
            trailing = ""
            while token and token[-1] in ".,;):":
                trailing = token[-1] + trailing
                token = token[:-1]
            add_hyperlink(paragraph, token, token, size=size)
            if trailing:
                set_run_font(paragraph.add_run(trailing), size=size, bold=bold, italic=italic)
        else:
            base = paragraph.add_run(match.group(1))
            set_run_font(base, size=size, bold=bold, italic=True, name="Cambria Math")
            subscript = paragraph.add_run(match.group(2))
            set_run_font(subscript, size=size, bold=bold, italic=True, name="Cambria Math")
            subscript.font.subscript = True
        cursor = match.end()
    if cursor < len(text):
        set_run_font(paragraph.add_run(text[cursor:]), size=size, bold=bold, italic=italic)


def set_front(document: Document, title: str, author: str, affiliations: list[str], email: str):
    paragraphs = document.paragraphs
    for paragraph in paragraphs[:10]:
        clear_paragraph(paragraph)
    title_parts = title.split(":", 1)
    for idx, text in ((0, title_parts[0]), (1, title_parts[1].strip() if len(title_parts) > 1 else "")):
        paragraphs[idx].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraphs[idx].add_run(text)
        set_run_font(run, size=24, bold=True)
    paragraphs[3].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(paragraphs[3].add_run(author), size=14)
    for idx, text in zip((5, 6), affiliations[:2]):
        paragraphs[idx].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(paragraphs[idx].add_run(text), size=10, italic=True)
    paragraphs[7].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(paragraphs[7].add_run(email), size=10, italic=True)


def add_body(document: Document, text: str, *, first_indent=True, italic=False, bold=False):
    paragraph = document.add_paragraph(style="ICEST_Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(0.5) if first_indent else Cm(0)
    add_formatted_text(paragraph, text, size=11, italic=italic, bold=bold)
    return paragraph


def add_heading(document: Document, text: str, level=1):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(6 if level == 1 else 3)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(paragraph.add_run(text), size=11 if level == 1 else 10, bold=level == 1, italic=level > 1)
    return paragraph


def add_abstract(document: Document, abstract: str, keywords: str, labels: tuple[str, str]):
    for label, text in ((labels[0], abstract), (labels[1], keywords)):
        paragraph = document.add_paragraph(style="ICEST_Abstract")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(0.5)
        run = paragraph.add_run(label + " - ")
        set_run_font(run, size=10, bold=True, italic=True)
        set_run_font(paragraph.add_run(text), size=10)


def add_equation(document: Document, equation: str, number: int):
    """Add a numbered equation prepared for conversion to native Word OMML.

    The [[OMML]] marker is consumed by convert_word_equations.ps1. A separate
    number paragraph keeps narrow two-column layouts stable after BuildUp().
    """
    paragraph = document.add_paragraph(style="ICEST_Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(f"[[OMML]]{equation}"), size=9, name="Cambria Math")
    number_paragraph = document.add_paragraph(style="ICEST_Normal")
    number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    number_paragraph.paragraph_format.first_line_indent = Cm(0)
    number_paragraph.paragraph_format.space_after = Pt(1)
    number_paragraph.paragraph_format.line_spacing = 0.8
    number_paragraph.paragraph_format.keep_with_next = True
    set_run_font(number_paragraph.add_run(f"({number})"), size=8)


def add_algorithm(document: Document, number: int, title: str, inputs: str, steps: list[str], output: str, lang: str):
    table = document.add_table(rows=1, cols=1)
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table, color="7F8C8D", size="6")
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    cell = table.cell(0, 0)
    set_cell_margins(cell, top=100, start=110, bottom=100, end=110)
    shade_cell(cell, "F3F6F8")
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    label = f"Algorithm {number}. {title}" if lang == "en" else f"Алгоритм {number}. {title}"
    set_run_font(paragraph.add_run(label), size=9, bold=True)
    p = cell.add_paragraph()
    set_run_font(p.add_run("Input: " if lang == "en" else "Вход: "), size=8, bold=True)
    set_run_font(p.add_run(inputs), size=8)
    for index, step in enumerate(steps, start=1):
        p = cell.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.25)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(f"{index}. {step}"), size=8, name="Consolas")
    p = cell.add_paragraph()
    set_run_font(p.add_run("Output: " if lang == "en" else "Выход: "), size=8, bold=True)
    set_run_font(p.add_run(output), size=8)


def add_equation_explanation(document: Document, text: str):
    paragraph = document.add_paragraph(style="ICEST_Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_after = Pt(2)
    add_formatted_text(paragraph, text, size=9, italic=True)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="808080", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def add_table(document: Document, headers: list[str], rows: list[list[str]], caption: str):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, "D9EAF7")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        clear_paragraph(paragraph)
        set_run_font(paragraph.add_run(header), size=8, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index else WD_ALIGN_PARAGRAPH.LEFT
            clear_paragraph(paragraph)
            set_run_font(paragraph.add_run(str(value)), size=8)
    paragraph = document.add_paragraph(style="ICEST_Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    set_run_font(paragraph.add_run(caption), size=10, italic=True)


def add_figure_prompt(document: Document, number: int, caption: str, prompt: str, lang: str):
    placeholder = document.add_table(rows=1, cols=1)
    placeholder.style = "Normal Table"
    placeholder.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(placeholder, color="A6A6A6", size="6")
    cell = placeholder.cell(0, 0)
    shade_cell(cell, "F2F2F2")
    set_cell_margins(cell, top=180, start=160, bottom=180, end=160)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_paragraph(paragraph)
    is_screenshot = number >= 7
    if lang == "en":
        label = f"Figure {number}: insert verified application screenshots" if is_screenshot else f"Figure {number} placeholder"
    else:
        label = f"Рисунок {number}: вставить проверенные скриншоты приложения" if is_screenshot else f"Место для рисунка {number}"
    set_run_font(paragraph.add_run(label), size=10, bold=True)
    cap = document.add_paragraph(style="ICEST_Normal")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    set_run_font(cap.add_run(f"Figure {number}. {caption}" if lang == "en" else f"Рисунок {number}. {caption}"), size=10, italic=True)
    prompt_p = document.add_paragraph(style="ICEST_Normal")
    prompt_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    prompt_p.paragraph_format.first_line_indent = Cm(0)
    if is_screenshot:
        prefix = (
            "Screenshot assembly instructions (remove before submission): "
            if lang == "en"
            else "Инструкция по сборке скриншотов (удалить перед подачей): "
        )
    else:
        prefix = (
            "Draft image-generation prompt (remove before submission): "
            if lang == "en"
            else "Промпт для генерации изображения (удалить перед подачей): "
        )
    set_run_font(prompt_p.add_run(prefix), size=8, bold=True)
    add_formatted_text(prompt_p, prompt, size=8, italic=True)


def add_bullets(document: Document, items: list[str]):
    for item in items:
        paragraph = document.add_paragraph(style="ICEST_Normal")
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_run_font(paragraph.add_run("• "), size=11, bold=True)
        set_run_font(paragraph.add_run(item), size=11)


def pct(value):
    return f"{100 * float(value):.2f}%"


def cohort_table(metrics: dict, lang: str):
    no = metrics["without_goso"]
    yes = metrics["with_goso"]
    m0, m1 = no["metrics"], yes["metrics"]
    if lang == "en":
        headers = ["Metric", "International mode", "Kazakhstan requirements", "Difference"]
        labels = [
            ("Programmes", str(no["programme_count"]), str(yes["programme_count"]), "0"),
            ("Raw EPVO provenance", pct(m0["epvo_provenance"]), pct(m1["epvo_provenance"]), "+38.64 pp"),
            ("Profile provenance, regulatory block excluded", pct(m0["epvo_provenance_excluding_regulatory"]), pct(m1["epvo_provenance_excluding_regulatory"]), "0.00 pp"),
            ("Raw semester alignment, ±1", pct(m0["semester_alignment_pm1"]), pct(m1["semester_alignment_pm1"]), "+13.75 pp"),
            ("Prerequisite-adjusted alignment", pct(m0["semester_alignment_prereq_adjusted_pm1"]), pct(m1["semester_alignment_prereq_adjusted_pm1"]), "+9.45 pp"),
            ("Semantic-adjusted alignment", pct(m0["semester_alignment_semantic_adjusted_pm1"]), pct(m1["semester_alignment_semantic_adjusted_pm1"]), "+5.61 pp"),
            ("International checklist", pct(m0["international_score"] / 100), pct(m1["international_score"] / 100), "0.00 pp"),
            ("Hard violations", str(int(m0["hard_violations"])), str(int(m1["hard_violations"])), "0"),
        ]
        caption = "Table 3. Fresh quality-eligible plans in international and Kazakhstan-regulated modes."
    else:
        headers = ["Метрика", "Международный режим", "Требования Казахстана", "Разница"]
        labels = [
            ("Программы", str(no["programme_count"]), str(yes["programme_count"]), "0"),
            ("Полное происхождение ЕПВО", pct(m0["epvo_provenance"]), pct(m1["epvo_provenance"]), "+38,64 п.п."),
            ("Профильное происхождение без нормативного блока", pct(m0["epvo_provenance_excluding_regulatory"]), pct(m1["epvo_provenance_excluding_regulatory"]), "0,00 п.п."),
            ("Сырое соответствие семестру, ±1", pct(m0["semester_alignment_pm1"]), pct(m1["semester_alignment_pm1"]), "+13,75 п.п."),
            ("С учётом пререквизитов", pct(m0["semester_alignment_prereq_adjusted_pm1"]), pct(m1["semester_alignment_prereq_adjusted_pm1"]), "+9,45 п.п."),
            ("С учётом педагогической сложности", pct(m0["semester_alignment_semantic_adjusted_pm1"]), pct(m1["semester_alignment_semantic_adjusted_pm1"]), "+5,61 п.п."),
            ("Международный чек-лист", pct(m0["international_score"] / 100), pct(m1["international_score"] / 100), "0,00 п.п."),
            ("Жёсткие нарушения", str(int(m0["hard_violations"])), str(int(m1["hard_violations"])), "0"),
        ]
        caption = "Таблица 3. Сравнение свежих планов в международном режиме и по требованиям Казахстана."
    return headers, [list(row) for row in labels], caption


EN = {
    "title": "Curriculum-KAG: Evidence-Constrained Curriculum Synthesis with National Expert Data, Semantic Retrieval, and Semester-Aware Planning",
    "abstract": (
        "Universities still design curricula manually by translating a programme goal into learning outcomes, selecting courses, checking credits and prerequisites, and satisfying regulatory rules. "
        "We present Curriculum-KAG, a fast decision-support system for generating classical or interdisciplinary bachelor's, master's, and doctoral curricula. "
        "It combines ten years of expert evidence from Kazakhstan's state higher-education platform, multilingual SBERT/BM25 retrieval, a curriculum knowledge graph, exact-credit optimisation, and independent constraint verification. "
        "The system produces one or three explainable semester plans in international or Kazakhstan-regulated mode while preserving expert control. Experiments demonstrate improved course-LO matching, feasible credits, forward prerequisites, and high semester alignment."
    ),
    "keywords": "knowledge-augmented generation, curriculum synthesis, learning outcomes, multilingual SBERT, constrained optimisation",
    "sections": [
        ("1. Introduction", [
            "Designing an educational programme manually is a multi-stage academic process. A programme team first defines the professional profile and purpose of the programme. It then formulates learning outcomes (LOs): concrete statements of what a graduate must know, understand, and be able to perform. Courses are selected to teach and assess these outcomes, credits are distributed across semesters, prerequisites are ordered, and the complete structure is checked against education-level, institutional, and regulatory requirements. A course that does not support any programme LO has no defensible place in an outcome-based curriculum [1], [2].",
            "This work uses the national government database of educational programmes of Kazakhstan, maintained within the state Unified Higher Education Platform. After the first full explanation, the article refers to this platform by its established Russian abbreviation EPVO. For approximately ten years, universities have submitted already designed bachelor's, master's, and doctoral programmes, while external specialists have evaluated the declared course-to-LO relationships. The resulting longitudinal corpus contains not only course titles but also real programme structures and accumulated expert judgement.",
            "Curriculum-KAG turns this evidence into a fast programme-design assistant. The user selects classical or interdisciplinary mode, education level, one or two professional fields, language, duration, credits, programme goal, and LOs. The system retrieves real courses, estimates course-to-LO relationships, transfers comparable expert evidence, builds a prerequisite graph, distributes courses across semesters, and returns one or three explainable alternatives. A deterministic verifier checks credits, workload, prerequisite direction, level, domain relevance, local regulatory components, and real LO coverage before a plan can become active.",
            "The research question is whether multilingual machine learning, national expert memory, a curriculum knowledge graph, and formal constraints can be combined to produce complete plans that are both useful to a university and independently verifiable. The principal contribution is not unconstrained text generation; it is an auditable synthesis method in which every course has a source, an LO rationale, a semester reason, and a machine-checkable role in the programme."
        ]),
        ("1.2. Research Hypotheses and Contributions", [
            "Four hypotheses are tested. H1: historical expert evaluations contain a learnable signal, so a model trained on EPVO can distinguish academically supported course-LO pairs better than an unadapted multilingual encoder. H2: when several courses are semantically plausible, expert evidence from programmes of the same level and field improves their ranking. H3: a constrained planner can transform probabilistic recommendations into complete curricula with the required credits, balanced semesters, forward prerequisites, and no hard violations. H4: Kazakhstan-regulated and international programmes must be evaluated separately because protected local components reduce scheduling freedom but should not be counted as irrelevant profile courses.",
            "The first contribution is a reproducible national dataset pipeline in which raw records, normalized entities, approved courses, and generation candidates are separated. The second is a two-stage recommendation architecture: a semantic classifier retrieves plausible links and an expert-memory reranker reorders the shortlist using historical evidence. The third is a graph-aware planner that treats credits, levels, domains, prerequisites, and statutory blocks as constraints rather than after-the-fact warnings.",
            "The fourth contribution is an explanation and feedback protocol. Every course can be traced to its source and linked outcomes; model evidence and expert evidence remain distinguishable; uncertain bridge modules are visible; and expert decisions are versioned. The fifth contribution is an independent validation layer that evaluates the stored plan after generation instead of trusting an internal optimisation score.",
            "Unlike a system that asks a language model to write a plausible list of course names, Curriculum-KAG operates over a bounded approved repository and preserves whole-course credits. Generative AI is used only in explicitly marked assistance functions. The core curriculum remains reproducible when external APIs are unavailable, a requirement for institutional use and scientific replication. The method extends the author's earlier work on digital educational platforms [17], AI-assisted syllabus design [18], graph and sequence models for curricula [19], and the initial Curriculum-KAG architecture [20]. The prospective evaluation layer can also draw on hybrid predictive modelling of educational attainment [21]."
        ]),
        ("2. Related Work", [
            "Sentence-BERT maps sentences to a shared vector space and enables efficient semantic retrieval [3]. Multilingual sentence representations are particularly important for the Kazakhstan context because course titles and descriptions appear in Kazakh, Russian, and English. BM25 complements dense retrieval by preserving exact technical expressions that may be diluted in a sentence embedding [4]. Knowledge-augmented generation extends retrieval by attaching structured entities, provenance, and rules to retrieved text [5].",
            "Curriculum construction also resembles multi-objective optimisation. Coverage, relevance, diversity, and interdisciplinarity must be maximised while credit deviation, duplication, prerequisite violations, and domain imbalance are minimised. NSGA-II provides a useful experimental Pareto-search branch [6], but the operational path in this study is a deterministic constrained planner with repair heuristics. This distinction improves reproducibility: the same input can be regenerated and independently verified, while evolutionary search remains available for controlled comparison.",
            "Graph neural networks and recurrent architectures can represent curricular dependencies, but higher complexity is not automatically beneficial. Controlled GNN and LSTM pilots did not outperform the semantic baseline on the available frozen splits. The production pipeline therefore retains SBERT-based retrieval and uses graph reasoning for prerequisite closure and verification rather than promoting an unconfirmed neural architecture [7]-[9]."
        ]),
        ("3. Data and Provenance", [
            "The source corpus was collected from the public Registry of Educational Programs within the Unified Higher Education Platform. The registry reflects a two-stage academic process. First, university departments formulate a programme and map courses to learning outcomes. Second, external experts review these mappings and assign low, medium, or high achievability. The recovered numerical values 0, 0.5, and 1 therefore represent expert relationship strength rather than neural-model confidence.",
            "The data pipeline preserves four layers: immutable raw EPVO records; normalized and deduplicated programmes, courses, outcomes, directions, and expert links; an approved multilingual course repository; and project-scoped generation candidates. This separation prevents raw imports from contaminating active curricula and makes every selected course traceable to its source. The operational PostgreSQL/pgvector repository stores three-language titles and descriptions, embeddings, prerequisite edges, model scores, expert feedback, and versioned plans.",
            "Programme-level train, validation, and test splits were used to reduce contextual leakage. All cards originating from one educational programme remain in the same split. Model thresholds and mixture weights are selected on validation data and then frozen before test evaluation. Dataset and model reports retain the seed, configuration, checksums, and split identifiers."
        ]),
        ("4. Method", [
            "For course c_i and learning outcome o_j, the multilingual encoder g_theta produces L2-normalised vectors. Their semantic similarity is the cosine of the angle between the vectors. A lexical score is added to protect exact professional terminology. The hybrid retrieval score is defined by Equation (1), where alpha is selected on validation data.",
            "Historical expert evidence is transferred conservatively. An EPVO score is used only when the source learning outcome and the new outcome share sufficient semantic or lexical content. Several expert votes are aggregated without collapsing medium and high relations into a single binary label. AI and EPVO scores are displayed separately because they answer different questions: the former is a model prediction for a new pair; the latter is evidence from previously reviewed programmes.",
            "A course is admitted only if it belongs to the selected education level and programme scope, supports at least one learning outcome, satisfies the domain guard, and can participate in a feasible prerequisite closure. Real EPVO courses are preferred. A bridge module is created only when no sufficiently supported real course closes an outcome or an exact credit gap, and it is visibly marked for expert confirmation.",
            "The planner assigns courses to semesters using the typical EPVO semester, prerequisite depth, semantic complexity, and semester capacity. Reranking is constrained: moving a course closer to its historical semester is allowed only if both source and destination loads remain inside the permitted interval and all prerequisite edges still point forward. This prevents a superficially better similarity score from producing an impossible plan."
        ]),
        ("4.1. Auditable Admission and Domain Evidence", [
            "Candidate admission is stricter than nearest-neighbour retrieval. Let I_level, I_scope, I_LO, and I_context denote binary checks for education level, EPVO direction or group, minimum learning-outcome evidence, and absence of a foreign professional context. A candidate can enter the optimisation pool only when their product equals one. This rule blocks, for example, a master's robotics course from entering a bachelor's programme and prevents a generic occurrence of the word digital from legitimising an unrelated professional discipline.",
            "Deduplication is performed before domain accounting. Several source cards may normalize to one canonical course while retaining evidence from different programmes. Domain evidence is therefore aggregated over every source record rather than inherited from the last database row. Exact programme-group evidence has priority over direction-level evidence, while a course supported by both selected domains can contribute to an interdisciplinary quota. The approach prevents database ordering from changing curriculum composition.",
            "A selected course must support at least one programme learning outcome. The final relationship is accepted through a two-threshold rule: strong model evidence can establish a new link, whereas sufficiently similar historical EPVO evidence can confirm it. Weak signals remain visible but do not satisfy the real-course coverage invariant. Protected SCES RK components receive dedicated system learning outcomes so that they appear in the OBE matrix without pretending to be professional EPVO evidence."
        ]),
        ("4.2. Exact Credits, Bridge Control, and Regulatory Closure", [
            "Credit fitting treats a real course as an atomic unit. Its original credit value cannot be silently reduced or split to close a semester gap. After the protected SCES RK block is inserted, the remaining professional-credit target is solved as an exact-fit selection problem over admissible course credits. Dynamic programming retains the best evidence-bearing combination for every reachable credit sum and selects the exact remainder before considering a synthetic filler.",
            "Bridge modules have two legitimate roles: they may represent genuinely new content at the intersection of fields, or they may close a learning-outcome gap for which the repository has no sufficiently supported real course. A generic component placeholder is never treated as a discipline. The replacement service retrieves three real alternatives, exposes title, description, credits, level, AI evidence, EPVO evidence, and prerequisites, and stores the expert's selection for the next transactional generation.",
            "Regulatory closure is conditional on programme mode and level. International programmes do not receive Kazakhstan-specific mandatory courses. A programme created under SCES RK receives only the applicable bachelor's, master's, or doctoral components, including practice, research work, and final attestation with their prescribed credits. These elements are protected from suspicious-course replacement and from diversity mutations across A/B/C."
        ]),
        ("4.3. Knowledge Graph, Semester Repair, and A/B/C Transactions", [
            "The curriculum graph contains course nodes, learning-outcome nodes, prerequisite edges, postrequisite relations, domain membership, and semester assignments. Explicit repository prerequisites are combined with conservative inferred edges based on foundation and advanced terminology. The graph is required to be acyclic after repair. For every edge c_a -> c_b, the verifier requires semester(c_a) < semester(c_b), except for separately classified concurrent regulatory activities.",
            "Semester repair uses bounded whole-course moves and swaps. It first addresses hard prerequisite or load violations, then improves pedagogical ordering and historical semester proximity. Moving a late advanced clinical, cybersecurity, research, or systems course into the first foundation term is prohibited by semantic floors. Introductory and foundation courses receive upper bounds so that one anomalous historical record cannot push them to the final year.",
            "Alternative generation varies ranking weights and candidate choices while preserving a common regulatory core. Pairwise fingerprints and Jaccard similarity ensure that A, B, and C are genuinely distinct. All requested alternatives are generated and verified before one atomic commit. This protects old active plans and makes a generation failure recoverable. Users can activate any successful alternative and retain the other versions for comparison."
        ]),
        ("4.4. Formal Objective and Independent Verifier", [
            "Let x_i indicate inclusion of course c_i and z_is its assignment to semester s. The planner maximises a weighted objective containing probabilistic LO coverage, domain relevance, expert evidence, diversity, and historical-semester agreement. Penalties represent duplicate content, credit deviation, quota deficits, bridge count, and prerequisite complexity. The weights define variants A, B, and C but do not relax hard constraints.",
            "The main constraints require each selected course to appear in exactly one semester, the total credits to equal the programme target within the declared tolerance, each semester to remain within its load interval, and every prerequisite to precede its dependent course. Domain credits must meet project quotas, and each professional LO must have at least one admitted real course unless an explicitly reviewed bridge is retained. Education-level and SCES RK guards are Boolean admission constraints.",
            "The verifier recomputes these properties from the persisted plan. Feasibility V(P) equals one only when the credit, load, prerequisite, level, duplicate, domain, LO-evidence, and regulatory predicates all equal one. This independent calculation prevents an optimisation routine from declaring success using stale intermediate state. International-quality indicators are reported alongside V(P) but cannot override a failed hard predicate.",
            "Coverage is probabilistic rather than a simple course count. If several moderately supported courses address one outcome, their combined evidence increases while remaining bounded by one. This avoids reporting 300% coverage and prevents a single weak course from automatically producing a full score. The interface also shows the contributing courses so that an expert can inspect whether statistical aggregation matches academic judgement."
        ]),
        ("5. Experimental Design", [
            "Evaluation separates pair classification, retrieval ranking, and complete-plan validity. ROC-AUC, PR-AUC, and F1 evaluate course-to-outcome classification. Recall@K measures the fraction of expert-relevant courses recovered in the first K positions, MRR rewards the position of the first relevant course, and nDCG@K evaluates the full graded ordering. At plan level, the verifier checks exact credits, semester load, prerequisites, education level, duplicate titles, domain quotas, real learning-outcome evidence, and protected regulatory components.",
            "During iterative functional testing, more than 80 classical and interdisciplinary programme configurations were successfully generated and reviewed across bachelor's, master's, and doctoral levels. For the reproducible numerical comparison reported here, six quality-eligible fresh programmes were selected as a balanced cohort: three operated in international mode and three followed the protected Kazakhstan regulatory mode. An additional deliberately infeasible stress case was excluded before aggregation. The same validator and tolerance of ±1 semester were used. This is structural validation, not a blinded accreditation study.",
            "SCES RK is the State Compulsory Standard of Education of the Republic of Kazakhstan. Depending on education level, it introduces a protected core of mandatory general education, practice, research work, and final attestation. These components cannot be replaced merely to increase semantic similarity. Consequently, raw EPVO provenance and semester flexibility are expected to be lower in the regulated cohort; profile-only metrics must therefore be reported separately."
        ]),
        ("6. Results", [
            "The base multilingual SBERT achieved ROC-AUC 0.7031, PR-AUC 0.6946, and F1 0.6886. Fine-tuning on 40,000 EPVO pairs improved these values to 0.7666, 0.7681, and 0.7188. A cached 12k candidate reached 0.7748, 0.7746, and 0.7268 but remains an experimental candidate rather than silently replacing the validated operational baseline. A separate ranking-loss model achieved Recall@10 0.7071, MRR 0.7815, and nDCG@10 0.6923 on its frozen benchmark.",
            "A second controlled retrieval experiment combined SBERT with global and programme-scoped expert memory. On its independent clean-v2 split, Recall@10 increased from 0.5348 to 0.5657, Recall@20 from 0.7476 to 0.7580, MRR from 0.2898 to 0.3740, and nDCG@10 from 0.3163 to 0.3735. The same result was preserved when the memory limit increased from 64 to 128 candidates, providing an initial robustness check. The production model was not changed by this experiment.",
            "The cohort comparison shows that plans in international mode achieved 86.84% semantic-adjusted semester alignment, whereas plans following Kazakhstan's protected requirements achieved 81.23%. The 5.61 percentage-point difference is attributed to lower scheduling freedom: the local-regulation cohort contains 35 protected components across three plans. The raw semester gap is larger, 13.75 points, but falls to 5.61 points after prerequisite and pedagogical adjustments. Both cohorts achieved 100% profile provenance after excluding the regulatory block, a 100% international checklist score, and zero hard violations.",
            "Fresh aggregation across all six eligible programmes produced raw semester agreement of 76.82%, prerequisite-adjusted agreement of 80.42%, and semantic-adjusted agreement of 84.03%. The constrained feasibility indicator is retained as a separate diagnostic and does not replace these historical-agreement metrics. On control project 137 it reached 100%, meaning that every remaining disagreement could not be repaired by a direct move without violating the 27-33-credit envelope; it does not mean that every historical semester was reproduced exactly."
        ]),
        ("6.1. End-to-End Control Programmes and Ablation Evidence", [
            "Control programmes cover conventional and interdisciplinary bachelor's designs, a master's programme, and a doctoral programme. Fresh generation produced exact totals of 240, 120, and 180 credits at the corresponding levels. The stored plans passed prerequisite, level, duplicate, credit-integrity, foreign-context, and regulatory checks. The international checklist reported 100% for all six quality-eligible programmes, while the predefined negative control remained outside the cohort rather than being used to inflate the mean.",
            "Ablation observations identify which safeguards cause the improvement. In the Digital Forensics Investigator case, the education-level guard removed courses from inappropriate levels, final LO repair established real coverage for every professional outcome, and the contextual guard removed unrelated courses hidden behind generic digital terminology. In IT-medicine, the real-course-first residual rule selected an admissible EPVO course before a generic credit filler and eliminated unnecessary bridges.",
            "The D094 doctoral control demonstrates exact fitting after a large regulatory block. The planner retained the protected research, practice, and attestation components and selected an exact professional remainder from real courses. Earlier bridge-heavy variants were replaced by evidence-bearing EPVO courses while maintaining 180 total credits and zero hard violations. This case is important because raw provenance alone would understate quality: most credits are intentionally regulatory.",
            "The A/B/C mechanism was also tested for diversity. Alternatives share mandatory components but differ in professional composition and ranking priorities. Transactional fingerprints prevent accidental publication of identical variants. When only one plan is requested, the system avoids the cost of building unused alternatives; A/B/C can later be completed from the same frozen project constraints."
        ]),
        ("6.2. Efficiency, Failure Handling, and Explainability", [
            "Repeated preprocessing was a major source of latency. Cache signatures now incorporate relevant data and configuration, so unchanged EPVO normalization and course-to-outcome scoring can be reused. In a controlled repeated build, the EPVO stage decreased from 28.45 to 0.15 seconds and the scoring stage from 175.38 to 0.16 seconds; full generation decreased from 320.5 to 138.3 seconds because the remaining time represented actual A/B/C synthesis and verification.",
            "External AI services are optional assistants for proposing programme goals, learning outcomes, achievability explanations, and bridge titles. Structured response contracts validate language, item count, score ranges, and domain relevance. Provider failure activates a deterministic fallback and never bypasses the curriculum verifier. Long operations expose progress rather than leaving the interface in an indefinite loading state.",
            "The explanation layer distinguishes final relationship strength, AI prediction, and EPVO evidence. A statement such as AI 76% and EPVO 21% means that two independent signals support the same course-LO pair; the numbers are not percentages of the programme and are not added. The interface also lists preceding and subsequent courses from the active plan, making the graph useful for curriculum review rather than merely decorative."
        ]),
        ("6.3. Registry Comparison and International Quality", [
            "A generated interdisciplinary programme is not expected to copy one historical programme. For each EPVO reference programme, the system calculates shared courses, Jaccard overlap, generated-plan containment, and reference recall. Low overlap with any single source can coexist with high provenance when the new programme intentionally combines validated components from several directions. The comparison screen therefore displays several neighbours and missing or distinctive elements instead of presenting one analogue as ground truth.",
            "The international checklist operationalises OBE, ABET-style continuous improvement, CDIO integration, and Tuning competence principles. It checks outcome coverage, credit and semester consistency, prerequisite integrity, domain relevance, interdisciplinary integration, assessment-method availability, and evidence of expert participation. The score is diagnostic: every component remains expandable and actionable, and a 100% software checklist is not described as international accreditation.",
            "The cohort results illustrate this distinction. Both groups obtained a 100% checklist score because their plans passed the implemented criteria. The semantic semester values remained below 100% because historical agreement is a separate external comparison. Maintaining both results prevents the system from hiding a semester discrepancy behind a broad quality score and gives the programme team a concrete target for manual or algorithmic refinement.",
            "The knowledge graph provides an additional visual audit. For a selected alternative, nodes are arranged by semester, prerequisite edges are directed, and the interface explains what the student is expected to acquire after each term. Course cards list prerequisites and postrequisites from the active plan. A future empirical study can compare these generated progression statements with student assessment data."
        ]),
        ("7. Discussion", [
            "The results show why regulatory and professional metrics must not be conflated. Raw provenance in the SCES RK cohort is 61.36% because protected practices, research work, and attestation are regulatory objects rather than ordinary EPVO courses. After excluding these expected objects, profile provenance is 100% in both cohorts. Reporting only the raw percentage would incorrectly portray statutory compliance as retrieval failure.",
            "The semester comparison has the same interpretation. International-mode programmes can move courses more freely toward historical EPVO positions, so they exceed the Kazakhstan-regulated cohort. Nevertheless, the local-requirement plans remain feasible, have no hard violations, and preserve all profile-course provenance. Curriculum-KAG therefore provides a strong initial plan rather than a final legal or accreditation decision.",
            "In a preliminary expert demonstration, specialists from Abylkas Saginov Karaganda Technical University and the Academy of Public Administration under the President of the Republic of Kazakhstan assessed the resulting programmes positively, particularly their traceability, explainability, and protection of regulatory components. They recommended limited manual refinement of disputed links and syllabus-level details before formal approval. This assessment supports the human-in-the-loop design: automation prepares a coherent evidence package, while the academic expert retains authority.",
            "The practical value lies in reducing repetitive work. A programme team can define the field, education level, goals, and learning outcomes; obtain scoped course candidates; inspect why each course was selected; compare alternatives; review prerequisite and postrequisite chains; and correct uncertain links. Expert actions are stored as future learning data, enabling continuous improvement without hiding human responsibility."
        ]),
        ("8. System Implementation and Reproducibility", [
            "The implemented system uses a FastAPI backend, a React interface, PostgreSQL as the primary transactional database, and pgvector for semantic indexing. Course, learning-outcome, prerequisite, expert-evidence, and plan-version entities are stored separately. This avoids embedding essential provenance in generated prose and permits the verifier to reproduce every decision directly from structured records. SQLite is retained only as a lightweight local fallback; experiments and the reported fresh generations use the PostgreSQL path.",
            "Generation is transactional. The planner constructs the requested alternative or the complete A/B/C set in memory, applies repair procedures, verifies each result, and commits only when the requested set satisfies the invariants. If a provider timeout, an impossible credit remainder, or a validation error occurs, the transaction is rolled back and the previous active plan remains unchanged. This property is important for academic information systems because a partial generation must never overwrite a previously approved curriculum.",
            "The planning process is divided into observable stages: scope selection, EPVO candidate retrieval, course-to-outcome scoring, regulatory closure, prerequisite closure, credit fitting, semester allocation, repair, and independent validation. The interface reports stage progress, prevents duplicate launches, and permits generation of either one plan or three alternatives. If one plan was initially requested, the remaining alternatives can be generated later without recreating the project definition.",
            "Course explanations expose the final relationship strength, the AI prediction, the transferred EPVO expert score, source direction and programme group, prerequisite role, and semester rationale. These numbers are not added as if they were parts of one percentage. They are separate evidence channels combined by an explicit policy. The expert may confirm a link, mark it as weak or incorrect, propose a correction, or exclude a course from the next regeneration. This feedback is versioned and can become supervised data in a later model cycle.",
            "The multilingual layer stores title and description records independently for Russian, Kazakh, and English, together with a source status. Verified source translations are distinguished from generated drafts requiring review. Interface labels, cycle names, errors, exports, and course explanations use the selected language with deterministic fallback. This design prevents a translated interface from silently displaying a course in a different language and preserves the origin of each text.",
            "Reproducibility artefacts include the frozen split, seed 42, dataset counts, model and configuration identifiers, metric JSON files, plan fingerprints, and automated regression tests. The current local suite covers semantic deduplication, education-level guards, regulatory protection, prerequisite order, credit tolerance, course uniqueness, multilingual integrity, and safe fallback when an external AI provider is unavailable. The same validator is called by the interface and by offline experiments, reducing divergence between a scientific benchmark and the operational application.",
            "The versioned source code, deployment instructions, experiment configurations, database migrations, and quality checks are available at https://github.com/kmypat-commits/curriculum-kag. Large model weights and the PostgreSQL dataset are distributed separately with checksums so that the repository remains practical to clone while each experiment can still be reconstructed."
        ]),
        ("9. Validity and Academic Use", [
            "Internal validation establishes software and structural validity: generated plans satisfy declared formal constraints, model experiments use programme-disjoint splits, and reported cohort metrics can be recomputed from saved artefacts. Content validity is supported by the longitudinal EPVO corpus and by the separation of AI prediction from historical expert evidence. Construct validity is strengthened by reporting classification, ranking, provenance, semester alignment, and whole-plan feasibility as different measurements rather than compressing them into one opaque score.",
            "More than 80 classical and interdisciplinary programme configurations were successfully generated and reviewed during iterative engineering validation. For a fully reproducible numerical comparison, the present study uses a frozen six-programme cohort: three programmes in international mode and three programmes with Kazakhstan's protected local requirements. This cohort supports a controlled engineering comparison but not a population-level claim about every discipline. Future evaluation can add a temporal holdout, blind review with a common rubric, inter-rater agreement, discipline-level error analysis, and prospective measurement of learning outcomes after programme implementation.",
            "For practical deployment, the system acts as a co-designer. It can prepare an initial programme structure, reveal missing outcome coverage, identify inconsistent prerequisites, compare the design with programmes in the Registry, and produce an auditable evidence package. The responsible academic body still approves the programme goal, learning outcomes, course contents, assessment strategy, staff and infrastructure requirements, and compliance with the current legal edition of SCES RK.",
            "The distinction between a bridge module and a real EPVO course is preserved throughout the lifecycle. A bridge represents a justified curriculum hypothesis at the intersection of fields. It becomes an ordinary approved course only after an expert selects or edits its title, description, credits, outcomes, prerequisites, and assessment approach. This workflow allows genuinely new interdisciplinary content without presenting synthetic modules as historical registry evidence.",
            "The next research stage will optimise semester ranking directly under credit and prerequisite constraints. Candidate moves will be evaluated globally rather than as isolated swaps, while the raw historical metric remains unchanged for comparability. Retrieval experiments will test programme-dependent hard negatives and listwise expert-strength objectives. A new model will be promoted only if it improves Recall, MRR, and nDCG on an independent split and preserves whole-plan invariants on fresh generation."
        ]),
        ("9.1. Research Roadmap", [
            "The first planned extension is a larger blind expert evaluation. Reviewers from several disciplinary groups will receive generated and reference fragments without knowing their origin and will score relevance, semester appropriateness, prerequisite coherence, and readiness for adoption. Inter-rater agreement and confidence intervals will complement the existing qualitative demonstration.",
            "The second extension is temporal validation. Programmes registered before a cut-off will form training memory, whereas later programmes will be held out. This will test whether the method predicts future academic choices rather than recalling contemporary duplicates. Direction-transfer experiments will measure performance in data-poor fields and determine when a bridge is preferable to a weak cross-domain match.",
            "The third extension is learning-to-rank with graded expert strength. Direct regression to 0, 0.5, and 1 did not improve retrieval in the pilot, so future work will use pairwise or listwise ordering within one programme. Hard negatives will be selected from the same level and nearby direction, making them semantically plausible but academically incorrect and therefore more informative than random negatives.",
            "Finally, the curriculum graph can support syllabus generation, weekly topic allocation, assessment mapping, and longitudinal programme monitoring. These functions should inherit the same provenance and verifier principles. A generated syllabus will be treated as a draft linked to credits and LOs, while measured student achievement and expert revisions will close the continuous-improvement loop."
        ]),
        ("10. Conclusion", [
            "Curriculum-KAG began as a method for designing interdisciplinary programmes at the intersection of two fields, such as IT and medicine or law and cybersecurity. During implementation, the same evidence and constraint architecture proved suitable for classical single-field programmes. The resulting system now supports bachelor's, master's, and doctoral levels, one or two professional directions, international mode, and the documented local requirements of Kazakhstan.",
            "In practical terms, a university enters the programme type, level, fields, goal, and learning outcomes. Curriculum-KAG then searches the state expert corpus, proposes real courses, explains which outcome each course supports, constructs prerequisite and postrequisite relations, balances credits by semester, and returns one or three alternatives. KAG is essential because retrieval is combined with a knowledge graph, provenance, expert evidence, and hard rules; a plausible but structurally invalid list cannot pass the verifier.",
            "The experiments show that multilingual semantic retrieval improves after domain adaptation, programme-scoped expert memory improves ranking, and the constrained planner preserves complete-plan feasibility. International-mode plans reached 86.84% semantic-adjusted semester alignment, while plans following Kazakhstan's protected requirements reached 81.23%; both retained 100% profile provenance and zero hard violations in the reported cohort. More than 80 additional programme configurations were successfully exercised during functional testing.",
            "The system is therefore best understood as a fast academic co-designer. It can substantially reduce routine search, mapping, and structural checking, but it does not replace a programme committee. Experts still approve the goal, learning outcomes, disputed relationships, bridge modules, assessment methods, staffing, infrastructure, and legal compliance. The remaining manual refinement is deliberately visible and recorded as feedback, making future versions more accurate and the final institutional decision auditable."
        ]),
    ],
}


RU = {
    "title": "Curriculum-KAG: Доказательный синтез учебных планов с экспертными данными, семантическим поиском и ограниченным распределением по семестрам",
    "abstract": (
        "Учебные планы по-прежнему создаются вручную: цель программы переводится в результаты обучения, затем подбираются дисциплины, кредиты, пререквизиты и нормативные компоненты. "
        "Представлена Curriculum-KAG — система быстрой генерации классических и междисциплинарных программ бакалавриата, магистратуры и докторантуры. "
        "Она объединяет десятилетнюю экспертную базу государственной платформы высшего образования Казахстана, многоязычный SBERT/BM25-поиск, граф знаний, точный подбор кредитов и независимую проверку ограничений. "
        "Система формирует один или три объяснимых семестровых плана в международном режиме либо по требованиям Казахстана, сохраняя экспертный контроль. Эксперименты подтвердили релевантность дисциплин, корректные кредиты, направленные пререквизиты и высокое соответствие семестрам."
    ),
    "keywords": "knowledge-augmented generation, синтез учебных планов, результаты обучения, многоязычный SBERT, ограниченная оптимизация",
    "sections": [
        ("1. Введение", [
            "Ручное проектирование образовательной программы состоит из нескольких академических этапов. Рабочая группа определяет профессиональный профиль и цель программы. Затем формулируются результаты обучения (РО) — конкретные утверждения о том, что выпускник должен знать, понимать и уметь выполнять. Под каждый РО выбираются дисциплины и методы оценивания, распределяются кредиты, выстраиваются пререквизиты и семестры, после чего структура проверяется по требованиям уровня образования, вуза и законодательства. Если дисциплина не обеспечивает ни одного РО, её включение в outcome-based программу не имеет достаточного обоснования [1], [2].",
            "Исследование использует государственную базу образовательных программ Казахстана, которая ведётся внутри Единой платформы высшего образования. После первого полного пояснения далее применяется принятое сокращение ЕПВО. Около десяти лет университеты загружали в неё уже разработанные программы бакалавриата, магистратуры и докторантуры, а внешние специалисты оценивали заявленные связи дисциплин с РО. Так сформирован уникальный продольный корпус: он содержит не только названия курсов, но и реальные структуры программ и накопленные экспертные мнения.",
            "Curriculum-KAG превращает этот корпус в инструмент быстрой разработки. Пользователь выбирает классический или междисциплинарный тип, уровень, одну или две профессиональные области, язык, длительность, кредиты, цель и РО. Система извлекает реальные дисциплины, оценивает связи дисциплина-РО, переносит сопоставимые экспертные доказательства, строит граф пререквизитов, распределяет курсы по семестрам и возвращает один либо три объяснимых варианта. До активации детерминированный verifier проверяет кредиты, нагрузку, порядок пререквизитов, уровень, предметную релевантность, локальные нормативные компоненты и реальное покрытие РО.",
            "Исследовательский вопрос состоит в том, можно ли объединить многоязычное машинное обучение, государственную экспертную память, граф знаний учебного плана и формальные ограничения так, чтобы университет получил практически полезный и независимо проверяемый результат. Научный вклад заключается не в свободной генерации текста, а в аудируемом синтезе: каждая дисциплина имеет источник, связь с РО, причину семестра и формально проверяемую роль в программе."
        ]),
        ("1.2. Гипотезы и научный вклад", [
            "Проверяются четыре ясные гипотезы. H1: исторические оценки содержат обучаемый сигнал, поэтому модель после обучения на ЕПВО лучше отличает академически подтверждённые пары дисциплина-РО, чем исходный многоязычный encoder. H2: если несколько дисциплин семантически подходят, экспертная память программ того же уровня и направления улучшает их порядок. H3: constrained-планировщик способен превратить вероятностные рекомендации в полный план с требуемыми кредитами, сбалансированными семестрами, направленными вперёд пререквизитами и без жёстких нарушений. H4: программы по местным требованиям Казахстана и международные программы необходимо оценивать отдельно, потому что защищённые нормативные компоненты уменьшают свободу расписания, но не должны считаться нерелевантными профильными дисциплинами.",
            "Первый вклад - воспроизводимый национальный конвейер данных с разделением raw, normalized, approved repository и generation candidates. Второй - двухэтапная рекомендация: семантический классификатор получает кандидатов, а экспертная память меняет их порядок. Третий - graph-aware планировщик, для которого кредиты, уровни, направления, пререквизиты и ГОСО являются ограничениями, а не предупреждениями после генерации.",
            "Четвёртый вклад - протокол объяснения и обратной связи. Каждая дисциплина прослеживается до источника и РО; прогноз модели отделён от экспертного доказательства; bridge видим; решение эксперта версионируется. Пятый вклад - независимый verifier сохранённого плана, который не доверяет внутреннему score оптимизатора.",
            "В отличие от запроса языковой модели на правдоподобный список названий, Curriculum-KAG работает в ограниченном утверждённом репозитории и сохраняет целые кредиты. Генеративный ИИ включается только в явно маркированные вспомогательные функции. Основной план воспроизводится и без внешнего API. Метод продолжает предыдущие работы автора по цифровым образовательным платформам [17], проектированию силлабусов с ИИ [18], графовым и последовательностным моделям учебных планов [19] и исходной архитектуре Curriculum-KAG [20]. Перспективная оценка результатов реализации программы может также использовать гибридное прогнозирование образовательных достижений [21]."
        ]),
        ("2. Связанные исследования", [
            "Sentence-BERT формирует векторные представления предложений и обеспечивает быстрый семантический поиск [3]. Для Казахстана особенно важна многоязычность, поскольку карточки дисциплин представлены на казахском, русском и английском языках. BM25 сохраняет точные профессиональные термины [4], а Knowledge-Augmented Generation дополняет поиск сущностями, происхождением и правилами [5].",
            "Проектирование плана является многокритериальной задачей: требуется максимизировать покрытие, релевантность и междисциплинарность при минимизации дублирования и нарушений. NSGA-II используется как экспериментальная Pareto-ветка [6]. Рабочий путь системы - детерминированный constrained-планировщик с repair-эвристиками, что обеспечивает повторяемость результата и независимую проверку.",
            "Контролируемые эксперименты GNN и LSTM не превзошли SBERT на зафиксированных выборках. Поэтому графовые правила применяются для замыкания пререквизитов и верификации, а усложнение нейронной архитектуры не внедряется без подтверждённого выигрыша [7]-[9]."
        ]),
        ("3. Данные и происхождение", [
            "Корпус получен из публичного Реестра образовательных программ ЕПВО. Данные отражают двухэтапную экспертизу. Сначала кафедра разрабатывает программу и связывает дисциплины с результатами обучения. Затем внешние эксперты оценивают достижимость связей как низкую, среднюю или высокую. Восстановленные значения 0, 0,5 и 1 являются силой экспертной связи, а не уверенностью нейронной модели.",
            "Конвейер разделён на четыре слоя: неизменяемые raw-данные; нормализованные и дедуплицированные программы, дисциплины, результаты и экспертные связи; утверждённый многоязычный репозиторий; кандидаты конкретного проекта. PostgreSQL/pgvector хранит три языковые версии, embeddings, пререквизиты, модельные оценки, обратную связь и версии планов.",
            "Разбиение train/validation/test выполнено на уровне программ: карточки одной университетской программы не попадают одновременно в обучение и тест. Пороги и веса выбираются на validation и фиксируются до теста. Для воспроизводимости сохраняются seed, конфигурация, контрольные суммы и идентификаторы split."
        ]),
        ("4. Метод", [
            "Для дисциплины c_i и результата обучения o_j многоязычный encoder g_theta формирует нормализованные векторы. Семантическая близость рассчитывается как косинус, а лексический сигнал сохраняет точные профессиональные термины. Итог поиска определяется формулой (1), где alpha выбирается по validation.",
            "Экспертный сигнал ЕПВО переносится консервативно: он используется только при достаточном совпадении исходного и нового результата обучения. Несколько экспертных голосов агрегируются без сведения средней и высокой связи к одной бинарной метке. В интерфейсе AI-score и EPVO-score показаны раздельно.",
            "Дисциплина допускается в план, если соответствует уровню и выбранному направлению, подтверждает хотя бы один результат обучения, проходит domain guard и допускает замыкание пререквизитов. Реальные дисциплины ЕПВО имеют приоритет. Bridge-модуль создаётся только при отсутствии сильной реальной замены и явно маркируется для экспертного решения.",
            "Распределение по семестрам учитывает типичный семестр ЕПВО, глубину пререквизитов, сложность и вместимость. Constrained-reranking переносит дисциплину ближе к историческому семестру лишь тогда, когда исходный и целевой семестры остаются в диапазоне нагрузки, а все стрелки пререквизитов направлены вперёд."
        ]),
        ("4.1. Аудируемый допуск и доменные доказательства", [
            "Допуск кандидата строже обычного nearest-neighbour поиска. Обозначим I_level, I_scope, I_LO и I_context бинарные проверки уровня образования, направления или группы ЕПВО, минимального доказательства РО и отсутствия чужого профессионального контекста. Дисциплина входит в пул оптимизации только при произведении индикаторов, равном единице. Поэтому магистерский курс робототехники не попадает в бакалавриат, а слово «цифровой» не легализует нерелевантную дисциплину.",
            "Дедупликация выполняется до учёта предметных областей. Несколько исходных карточек могут образовать одну каноническую дисциплину, сохранив доказательства разных программ. Доменные оценки агрегируются по всем источникам, а не берутся из последней строки базы. Точное совпадение группы ОП имеет приоритет над направлением; дисциплина с доказательствами обеих областей учитывается как междисциплинарная.",
            "Каждая выбранная дисциплина должна поддерживать хотя бы один РО программы. Двухпороговое правило допускает сильное новое доказательство модели или подтверждение достаточно похожей исторической связью ЕПВО. Слабые сигналы видимы эксперту, но не закрывают инвариант реального покрытия. Для защищённых компонентов ГОСО создаются отдельные системные РО без выдачи их за профильные данные ЕПВО."
        ]),
        ("4.2. Точные кредиты, bridge-контроль и нормативное замыкание", [
            "Реальная дисциплина является атомарной кредитной единицей: её кредиты нельзя незаметно уменьшить или разделить. После вставки защищённого блока ГОСО профильный остаток решается как exact-fit задача. Динамическое программирование сохраняет лучший набор доказательных дисциплин для каждой достижимой суммы и выбирает точный остаток до создания синтетического заполнителя.",
            "Bridge-модуль допустим как действительно новое содержание на стыке областей либо как закрытие РО, для которого нет сильной реальной дисциплины. Пустой «компонент по выбору» не считается дисциплиной. Сервис замены показывает три реальных варианта, их описание, кредиты, уровень, AI/EPVO evidence и пререквизиты; выбор эксперта сохраняется для следующей генерации.",
            "Нормативное замыкание зависит от режима и уровня. Международная программа не получает казахстанские обязательные дисциплины. Режим ГОСО добавляет только применимые компоненты бакалавриата, магистратуры или докторантуры, включая практики, НИР и итоговую аттестацию с заданными кредитами. Они защищены от удаления и случайной мутации между A/B/C."
        ]),
        ("4.3. Граф знаний, semester-repair и транзакция A/B/C", [
            "Граф содержит дисциплины, РО, пререквизиты, постреквизиты, предметные области и семестры. Явные связи репозитория дополняются осторожно выведенными зависимостями «основа - продвинутый курс». После repair граф обязан быть ацикличным. Для каждого ребра c_a -> c_b verifier требует semester(c_a) < semester(c_b), кроме отдельно классифицированных параллельных нормативных активностей.",
            "Semester-repair использует ограниченные переносы и обмены целых дисциплин. Сначала устраняются жёсткие нарушения, затем улучшается педагогический порядок и историческая близость. Продвинутый клинический, кибербезопасностный или исследовательский курс не переносится в первый фундаментальный семестр. Вводные дисциплины получают верхнюю границу, чтобы единичная аномальная карточка ЕПВО не отправляла их в выпускной год.",
            "A/B/C различаются весами ranking и выбором кандидатов при общем нормативном ядре. Fingerprint и pairwise Jaccard проверяют реальное различие. Все запрошенные варианты строятся и проверяются до единого atomic commit. При ошибке старые активные планы сохраняются, а пользователь может активировать любой успешный вариант."
        ]),
        ("4.4. Формальная целевая функция и независимый verifier", [
            "Пусть x_i показывает включение дисциплины c_i, а z_is - назначение в семестр s. Планировщик максимизирует взвешенную функцию вероятностного покрытия РО, доменной релевантности, экспертного доказательства, разнообразия и исторической близости семестра. Штрафуются дубли, кредитное отклонение, дефицит квот, число bridge и сложность графа. Разные веса формируют A/B/C, но не ослабляют жёсткие ограничения.",
            "Ограничения требуют назначения каждой выбранной дисциплины ровно в один семестр, соблюдения итоговых кредитов и допустимой нагрузки, предшествования каждого пререквизита, выполнения квот областей и реального доказательства каждого профессионального РО. Level guard и ГОСО являются булевыми условиями допуска.",
            "Verifier пересчитывает свойства из сохранённого плана. Feasibility V(P)=1 только при одновременном выполнении предикатов кредитов, нагрузки, пререквизитов, уровня, отсутствия дубликатов, квот, LO-evidence и ГОСО. Это не позволяет оптимизатору объявить успех по устаревшему промежуточному состоянию. Международный score не может перекрыть провал жёсткого условия.",
            "Покрытие РО вероятностное, а не простая сумма дисциплин. Несколько средних доказательств усиливают результат, оставаясь в диапазоне до единицы. Поэтому система не показывает 300% покрытия и не считает одну слабую связь полноценным достижением. Интерфейс раскрывает дисциплины-источники для академической проверки."
        ]),
        ("5. Дизайн эксперимента", [
            "Классификация связи оценивается ROC-AUC, PR-AUC и F1. Retrieval оценивается Recall@K, MRR и nDCG@K. На уровне полного плана проверяются точные кредиты, нагрузка, пререквизиты, уровень образования, дубликаты, квоты областей, реальное покрытие РО и защищённые нормативные компоненты.",
            "В ходе итерационного функционального тестирования успешно сформировано и просмотрено более 80 конфигураций классических и междисциплинарных программ бакалавриата, магистратуры и докторантуры. Для воспроизводимого численного сравнения в статье использована сбалансированная свежая выборка из шести quality-eligible программ: три работали в международном режиме, ещё три — по защищённым местным требованиям Казахстана. Отдельный намеренно невыполнимый stress-test исключён до агрегации. Для обеих групп применены один validator и допуск ±1 семестр."
            "ГОСО - государственный общеобязательный стандарт образования Республики Казахстан. Он задаёт защищённое нормативное ядро: обязательные компоненты, практики, научно-исследовательскую работу и итоговую аттестацию в зависимости от уровня образования. Их нельзя заменить только ради роста семантического процента. Поэтому полное происхождение ЕПВО и свобода перемещения по семестрам ожидаемо ниже, а профильные метрики показываются отдельно."
        ]),
        ("6. Результаты", [
            "Исходная multilingual SBERT получила ROC-AUC 0,7031, PR-AUC 0,6946 и F1 0,6886. Дообучение на 40 000 парах ЕПВО повысило результаты до 0,7666, 0,7681 и 0,7188. Cached-кандидат 12k достиг 0,7748, 0,7746 и 0,7268, но сохранён как экспериментальный кандидат. Модель ranking-loss получила Recall@10 0,7071, MRR 0,7815 и nDCG@10 0,6923.",
            "Комбинация SBERT с глобальной и scoped-памятью экспертов на independent clean-v2 split повысила Recall@10 с 0,5348 до 0,5657, Recall@20 с 0,7476 до 0,7580, MRR с 0,2898 до 0,3740 и nDCG@10 с 0,3163 до 0,3735. Результат сохранился при расширении памяти с 64 до 128 кандидатов. Рабочая production-модель этим экспериментом не заменялась.",
            "В международном режиме семантически скорректированное соответствие семестру достигло 86,84%, а в режиме местных требований Казахстана — 81,23%. Разница 5,61 процентного пункта связана с меньшей свободой планирования: в трёх нормативных программах присутствуют 35 защищённых компонентов. Сырая разница равна 13,75 пункта, но уменьшается после учёта пререквизитов и педагогической сложности. Обе группы получили 100% профильного происхождения из ЕПВО, 100% международного чек-листа и ноль жёстких нарушений.",
            "По всем шести пригодным программам сырое соответствие семестру составило 76,82%, с учётом пререквизитов - 80,42%, с учётом семантических границ - 84,03%. Constrained-feasibility сохраняется отдельной диагностикой. На проекте 137 она составила 100%: оставшиеся прямые переносы нарушали бы диапазон 27-33 кредита. Это не означает буквальное совпадение каждого исторического семестра."
        ]),
        ("6.1. Сквозные контрольные программы и ablation", [
            "Контроль включает обычные и междисциплинарные программы бакалавриата, магистратуру и докторантуру. Свежая генерация дала точные объёмы 240, 120 и 180 кредитов. Планы прошли проверки пререквизитов, уровня, дубликатов, целостности кредитов, чужого контекста и ГОСО. Международный чек-лист равен 100% для шести quality-eligible программ; заранее заданный отрицательный контроль не включён в среднее.",
            "Ablation показывает вклад защитных механизмов. В программе «Киберследователь» level guard удалил дисциплины чужого уровня, финальный LO-repair обеспечил реальное покрытие профессиональных РО, context guard убрал нерелевантные курсы со случайным цифровым термином. В IT-медицине real-course-first правило выбрало допустимую дисциплину ЕПВО вместо общего заполнителя и сократило bridge.",
            "Докторская программа D094 демонстрирует exact-fit после крупного нормативного блока. Планировщик сохранил практики, НИР и аттестацию, затем набрал точный профильный остаток реальными дисциплинами. Bridge-heavy варианты были заменены доказательными курсами при 180 кредитах и нуле жёстких нарушений. Сырая provenance здесь ожидаемо ниже из-за нормативных кредитов.",
            "Проверено и разнообразие A/B/C. Обязательные компоненты совпадают, но профильный состав и ranking-приоритеты различаются. Транзакционные fingerprints запрещают публикацию одинаковых альтернатив. При выборе одной программы система не тратит время на две неиспользуемые, но позволяет достроить их позднее из тех же ограничений проекта."
        ]),
        ("6.2. Производительность, отказоустойчивость и объяснимость", [
            "Повторная предобработка была главным источником задержки. Подпись кэша учитывает данные и конфигурацию, поэтому неизменившиеся этапы ЕПВО и scoring переиспользуются. В повторном контрольном запуске этап ЕПВО сократился с 28,45 до 0,15 секунды, scoring - со 175,38 до 0,16 секунды, а полная A/B/C генерация - с 320,5 до 138,3 секунды. Оставшееся время относится к реальному синтезу и verifier.",
            "Внешние AI API используются только для предложений целей, РО, объяснений достижимости и названий bridge. Типизированный контракт проверяет язык, число элементов, диапазон оценок и предметную релевантность. При сбое включается детерминированный fallback, который не обходит verifier. Длительные операции показывают этапы прогресса и блокируют повторный запуск.",
            "Объяснение разделяет итоговую силу связи, прогноз ИИ и доказательство ЕПВО. AI 76% и EPVO 21% означают два независимых сигнала для одной пары, а не доли программы и не числа для сложения. Интерфейс дополнительно показывает предшествующие и последующие дисциплины активного плана, превращая граф в инструмент экспертизы."
        ]),
        ("6.3. Сравнение с Реестром и международное качество", [
            "Новая междисциплинарная программа не обязана копировать один исторический план. Для каждой программы ЕПВО рассчитываются общие дисциплины, Jaccard, containment и recall эталона. Низкое совпадение с одной программой совместимо с высокой provenance, если новый план объединяет подтверждённые элементы нескольких направлений. Экран сравнения показывает несколько аналогов, отсутствующие и уникальные элементы.",
            "Международный чек-лист операционализирует OBE, ABET-style continuous improvement, CDIO и Tuning. Проверяются РО, кредиты, семестры, пререквизиты, предметная релевантность, междисциплинарность, методы оценивания и участие эксперта. Score является диагностикой: 100% программной проверки не называется международной аккредитацией.",
            "Обе когорты получили 100% чек-листа, поскольку прошли реализованные критерии. Семестровое соответствие осталось ниже 100%, так как это отдельное внешнее сравнение. Раздельные результаты не позволяют скрыть проблему семестра общим score и дают ясную цель для ручной или алгоритмической коррекции.",
            "Граф знаний обеспечивает визуальный аудит: узлы расположены по семестрам, стрелки пререквизитов направлены вперёд, а интерфейс объясняет ожидаемый итог каждого семестра. Карточка показывает пре- и постреквизиты активного плана. В будущем эти выводы можно сопоставить с фактическими оценками студентов."
        ]),
        ("7. Обсуждение", [
            "Результаты показывают, почему нормативные и профильные показатели нельзя смешивать. В группе местных требований полное происхождение ЕПВО равно 61,36%, поскольку практики, исследовательская работа и аттестация являются нормативными объектами, а не обычными карточками дисциплин ЕПВО. После их корректного исключения профильное происхождение равно 100% в обеих группах.",
            "У программ международного режима больше свободы для приближения к историческим семестрам, поэтому их процент выше. При этом планы по требованиям Казахстана остаются выполнимыми, не содержат жёстких нарушений и сохраняют профильные дисциплины ЕПВО. Система формирует сильный исходный вариант, но не заменяет юридическое утверждение и аккредитацию.",
            "В ходе предварительной экспертной демонстрации специалисты Карагандинского технического университета имени Абылкаса Сагинова и Академии государственного управления при Президенте Республики Казахстан высоко оценили связность, объяснимость и защиту нормативных компонентов. Перед формальным утверждением рекомендована небольшая ручная докрутка спорных связей и деталей силлабусов. Это соответствует human-in-the-loop архитектуре.",
            "Практическая ценность состоит в сокращении рутинной работы. Разработчик задаёт идею, направление, уровень, цели и результаты; получает релевантные дисциплины, причины выбора, альтернативы, граф пререквизитов и отчёт о качестве. Экспертные подтверждения сохраняются как данные для дальнейшего улучшения."
        ]),
        ("8. Реализация и воспроизводимость", [
            "Система реализована на FastAPI и React; основной транзакционной базой является PostgreSQL, а pgvector используется для семантического индекса. Дисциплины, результаты обучения, пререквизиты, экспертные доказательства и версии планов хранятся раздельно. Поэтому происхождение не растворяется в сгенерированном тексте, а verifier может повторно получить каждое решение из структурированных записей. SQLite сохраняется только как облегчённый локальный fallback; эксперименты и свежая генерация статьи выполнены в PostgreSQL-режиме.",
            "Генерация транзакционна. Планировщик формирует выбранный вариант или полный набор A/B/C в памяти, применяет repair-процедуры, проверяет каждый результат и фиксирует данные только после прохождения инвариантов. При таймауте провайдера, невозможном кредитном остатке или ошибке проверки транзакция откатывается, а предыдущий активный план сохраняется. Частичный результат не может заменить ранее утверждённую версию.",
            "Процесс разделён на наблюдаемые этапы: выбор scope, получение кандидатов ЕПВО, оценка дисциплина-РО, нормативное замыкание, замыкание пререквизитов, подбор точной суммы кредитов, распределение по семестрам, repair и независимая проверка. Интерфейс показывает прогресс и блокирует повторный запуск. Пользователь может создать один вариант либо сразу A/B/C, а недостающие альтернативы построить позднее.",
            "Объяснение дисциплины показывает итоговую силу связи, прогноз ИИ, перенесённую оценку экспертов ЕПВО, направление и группу ОП, роль в графе и причину семестра. Эти проценты не складываются: они являются отдельными каналами доказательств. Эксперт может подтвердить связь, отметить её как слабую или неверную, исправить или исключить дисциплину из следующей генерации. Решение версионируется и может стать обучающим примером.",
            "Многоязычный слой отдельно хранит названия и описания на русском, казахском и английском языках вместе со статусом источника. Проверенные переводы отделены от черновиков, требующих просмотра. Подписи интерфейса, циклы, ошибки, экспорт и объяснения выбираются по текущему языку с детерминированным fallback. Это предотвращает незаметное смешение языков.",
            "Артефакты воспроизводимости включают frozen split, seed 42, статистику набора, идентификаторы модели и конфигурации, JSON с метриками, fingerprints планов и автоматические regression-тесты. Тесты покрывают дедупликацию, level guard, защиту ГОСО, порядок пререквизитов, допуск кредитов, уникальность дисциплин, целостность трёх языков и безопасный fallback при недоступности внешнего AI API. Один validator используется в интерфейсе и экспериментах.",
            "Версионированный исходный код, инструкция развёртывания, конфигурации экспериментов, миграции базы и проверки качества доступны по адресу https://github.com/kmypat-commits/curriculum-kag. Крупные веса моделей и набор PostgreSQL распространяются отдельно с контрольными суммами, чтобы репозиторий оставался пригодным для клонирования, а эксперимент можно было восстановить."
        ]),
        ("9. Валидность и академическое применение", [
            "Внутренняя проверка подтверждает программную и структурную валидность: планы соблюдают формальные ограничения, модельные эксперименты используют program-disjoint split, а когортные метрики пересчитываются из сохранённых файлов. Содержательная валидность поддерживается многолетним корпусом ЕПВО и разделением прогноза ИИ и исторического экспертного сигнала. Разные аспекты - классификация, ranking, provenance, семестр и выполнимость - не сворачиваются в один непрозрачный балл.",
            "Сравнение интерпретируется в наблюдаемом масштабе. Более 80 успешно просмотренных конфигураций подтверждают функциональную устойчивость, а сбалансированная выборка из трёх международных и трёх нормативных программ обеспечивает воспроизводимое численное сравнение. Она не является заявлением обо всех направлениях Казахстана. Для аудита сохранены идентификаторы программ, отрицательный контроль и правило формирования когорт. Следующий этап может включать temporal holdout, слепую экспертизу по общей рубрике, inter-rater agreement и перспективное измерение фактических результатов студентов."
            "В эксплуатации система выступает академическим соавтором решения: готовит структуру, показывает пробелы РО, находит конфликты пререквизитов, сравнивает с Реестром и формирует аудируемый пакет доказательств. Ответственный коллегиальный орган утверждает цель, РО, содержание, методы оценивания, кадровые и инфраструктурные требования и соответствие актуальной редакции ГОСО.",
            "Bridge-модуль не смешивается с реальной дисциплиной ЕПВО. Он является обоснованной гипотезой на стыке областей и становится утверждённой дисциплиной только после выбора или редактирования названия, описания, кредитов, РО, пререквизитов и оценивания экспертом. Так система поддерживает новую междисциплинарность, не выдавая синтетический модуль за исторический факт.",
            "Следующий исследовательский этап будет оптимизировать semester ranking непосредственно под ограничениями кредитов и графа. Переносы оцениваются глобально, а сырая историческая метрика сохраняется для сопоставимости. Retrieval-эксперименты проверят программно-зависимые hard negatives и listwise-цель с экспертной силой. Новая модель будет внедрена только при росте Recall, MRR и nDCG на независимом split и сохранении инвариантов свежей генерации."
        ]),
        ("9.1. Исследовательская дорожная карта", [
            "Первое расширение - более крупная слепая экспертиза. Специалисты нескольких направлений получат фрагменты сгенерированных и реальных планов без указания источника и оценят релевантность, семестр, пререквизиты и готовность к применению. Inter-rater agreement и доверительные интервалы дополнят качественную демонстрацию.",
            "Второе расширение - temporal validation. Программы до даты отсечения формируют память, более поздние остаются в holdout. Это проверит прогноз будущих академических решений, а не запоминание современных дублей. Domain-transfer покажет поведение в малых направлениях и границу выбора bridge вместо слабого междоменного кандидата.",
            "Третье расширение - learning-to-rank с градуированной экспертной силой. Прямая регрессия к 0, 0,5 и 1 не улучшила retrieval в пилоте, поэтому будут проверены pairwise/listwise цели внутри программы. Hard negatives выбираются из того же уровня и близкого направления: они семантически правдоподобны, но академически неверны.",
            "Наконец, граф поддержит генерацию силлабусов, недельных тем, assessment mapping и мониторинг программы. Эти функции наследуют provenance и verifier. Силлабус остаётся draft, связанным с кредитами и РО, а фактическое достижение студентов и экспертные исправления замыкают continuous-improvement цикл."
        ]),
        ("10. Заключение", [
            "Curriculum-KAG первоначально задумывалась для проектирования междисциплинарных программ на стыке двух областей, например IT и медицины либо права и кибербезопасности. В ходе реализации оказалось, что тот же доказательный и ограничительный конвейер подходит для классических программ одного направления. Итоговая система поддерживает бакалавриат, магистратуру и докторантуру, одну или две профессиональные области, международный режим и документированные местные требования Казахстана.",
            "На практике университет вводит тип и уровень программы, направления, цель и результаты обучения. Curriculum-KAG ищет реальные дисциплины в государственной экспертной базе, объясняет, какой РО обеспечивает каждая дисциплина, строит пре- и постреквизиты, балансирует кредиты по семестрам и выдаёт один либо три варианта. KAG принципиален, потому что поиск дополнен графом знаний, происхождением, экспертными оценками и жёсткими правилами: правдоподобный, но структурно неверный список не проходит verifier.",
            "Эксперименты показали улучшение многоязычного семантического поиска после обучения на предметных данных, рост качества ranking при подключении scoped-памяти экспертов и сохранение выполнимости полного плана. Международный режим достиг 86,84% семантически скорректированного соответствия семестрам, режим требований Казахстана — 81,23%; обе группы сохранили 100% профильного происхождения и ноль жёстких нарушений. В ходе функционального тестирования дополнительно успешно проверено более 80 конфигураций программ.",
            "Поэтому систему следует понимать как быстрого академического соавтора. Она сокращает ручной поиск, сопоставление и структурную проверку, но не заменяет учебно-методический орган. Эксперты по-прежнему утверждают цель, РО, спорные связи, bridge-модули, методы оценивания, кадровые и инфраструктурные условия и юридическое соответствие. Оставшаяся ручная доработка намеренно видима и сохраняется как обратная связь, поэтому итоговое решение можно проверить, а следующие версии — улучшать."
        ]),
    ],
}


SCIENTIFIC_INSERTIONS = {
    "en": {
        "1. Introduction": [
            ("1.1. What EPVO Is and Why It Matters", [
                "EPVO is the established Russian-language abbreviation for Kazakhstan's state Unified Higher Education Platform. Its Registry of Educational Programs is the government information system in which universities submit programmes for review and registration. For approximately ten years, this process has accumulated a unique longitudinal database of already designed programmes and external expert judgements. EPVO is therefore not a generic course catalogue: it contains programme identity, education level, field and group codes, goals, learning outcomes, courses, credits, recommended semesters, descriptions, and course-to-outcome mappings.",
                "The evidence is produced through an academic lifecycle. A university department first defines a programme goal and intended learning outcomes, constructs the curriculum, and declares which courses contribute to which outcomes. External reviewers working in the corresponding field then inspect the submitted programme. For a declared course-to-outcome relation they may assign low, medium, or high achievability. In the recovered data these ordinal judgements are encoded as 0, 0.5, and 1. Multiple votes are retained and aggregated rather than converted to a binary yes/no label.",
                "A simple example clarifies the distinction. If a department links Digital Forensics to an outcome on collecting and preserving digital evidence, SBERT can estimate a new semantic probability from the two texts. EPVO evidence answers a different question: how strongly did reviewers support comparable links in previously submitted programmes? Curriculum-KAG keeps these signals separate, displays both, and admits the course only after education-level, field, credit, prerequisite, and context checks also pass.",
                "This provenance makes the corpus more valuable than an unreviewed list of course names, but it does not make every historical decision universally correct. Programmes were created by different institutions, at different times, and for different local contexts. The system therefore treats EPVO as expert memory and evidence, not as an oracle. Programme-disjoint evaluation, conservative transfer, independent verification, and human confirmation are used to control memorisation and inappropriate reuse.",
                "The abbreviation SCES RK denotes the State Compulsory Educational Standard of the Republic of Kazakhstan; the Russian abbreviation GOSO RK is also widely used. It specifies protected components for the relevant education level. These components are structurally necessary in the Kazakhstan-compliant mode even when they are not ordinary profile courses in EPVO. International and SCES-constrained programmes are consequently evaluated with both complete-plan and profile-only denominators."
            ])
        ],
        "2. Related Work": [
            ("2.1. Why the System Is Called KAG", [
                "Retrieval-Augmented Generation (RAG) usually retrieves relevant text fragments and provides them as context to a generative model. This improves factual grounding, but the retrieved passages do not by themselves express which course precedes another, which learning outcome is supported, which expert supplied the evidence, or whether the complete plan satisfies credits and regulation. RAG can answer a question from documents; it does not automatically construct a valid curriculum.",
                "Knowledge-Augmented Generation (KAG) adds structured knowledge to retrieval. In Curriculum-KAG, the knowledge layer contains canonical courses, programme levels and fields, graded expert course-LO links, prerequisite and postrequisite edges, domain membership, regulatory status, credit values, semester evidence, and provenance. Retrieved candidates are expanded and filtered through this graph before the constrained planner acts. The final verifier then checks the complete stored plan.",
                "KAG is therefore better suited than plain RAG to this specific engineering task, not universally better for every language application. The advantage is relational and operational: the system can explain why a course was selected, identify what must be studied before and after it, distinguish a real EPVO course from a synthetic bridge, and reject an attractive textual answer that violates academic constraints. This article develops the earlier Curriculum-KAG concept [20] into a universal generator for both classical and two-domain interdisciplinary programmes."
            ])
        ],
        "3. Data and Provenance": [
            ("3.1. Formal Data Model", [
                "A programme p is represented by its education level h_p, field d_p, educational-program group b_p, goal text, learning-outcome set O_p, and original course set C_p. A raw course record stores a title, multilingual description, credits, component type, typical semester, and declared links to outcomes. A normalized course is a canonical entity connected to one or more source records. This many-to-one relation preserves provenance while preventing duplicate titles from occupying multiple places in a generated plan.",
                "An expert observation is a tuple (p,c,l,e), where course c was evaluated against source learning outcome l in programme p and e belongs to {0,0.5,1}. The same canonical course may therefore receive different judgements in different programmes. The method never replaces this contextual distribution with a single global truth; it uses a programme-scoped aggregate first and a broader aggregate only as fallback.",
                "The operational repository is produced through the sequence raw EPVO -> normalized EPVO -> approved multilingual repository -> project-scoped candidates. Raw objects and checksums are immutable. Normalization standardises whitespace, codes, credit formats, and multilingual fingerprints. Approval establishes a usable canonical record. Project scoping then filters by level, selected direction and group before semantic ranking begins.",
                "This layered design supports a dataset passport. The passport records extraction time, source checksum, schema version, language completeness, duplicates, label distribution, split assignment, and known limitations. The model benchmark records model identifier, seed, threshold, configuration checksum, and the exact programme identifiers used in train, validation, and test sets."
            ]),
            ("3.2. Expert-Label Reconstruction and Leakage Control", [
                "The recovered external judgements were initially at risk of being collapsed into a binary relation during normalization. The corrected pipeline preserves the graded values. In total, 876,842 source votes were aggregated into 800,793 contextual course-LO pairs; these form an evidence feature and an evaluation target, not an automatic curriculum decision.",
                "Cards from the same educational programme are never split across training and testing. This programme-disjoint rule is stricter than random pair splitting because courses and outcomes inside one programme share vocabulary and structure. Without grouping, near-duplicate context could produce an optimistic estimate of generalisation. Thresholds, mixture weights, and early stopping are chosen on validation programmes and frozen before the test programmes are scored.",
                "Negative examples are sampled within the same education level and a related field whenever possible. Such examples are harder than random cross-domain pairs and better approximate the practical ranking task: the system must distinguish several plausible courses, not merely separate medicine from an unrelated engineering outcome. Experimental programme-scoped memory and ranking-loss branches are evaluated separately from the operational classifier."
            ])
        ],
        "4.4. Formal Objective and Independent Verifier": [
            ("4.5. Optimisation Variables, Guarantees, and Complexity", [
                "Let x_i be a binary course-selection variable and z_is a binary assignment of course i to semester s. The admissible candidate set is fixed before optimisation by education-level, field, group, context, and minimum-evidence guards. The planner therefore optimises only over academically plausible candidates; it cannot compensate for an inadmissible course by giving it a favourable objective weight.",
                "The objective combines bounded LO coverage, semantic and expert relevance, domain balance, diversity, semester agreement, and bridge minimisation. Hard constraints are kept outside the weighted sum. This separation is important: no increase in similarity may purchase a prerequisite violation, an incorrect education level, or missing credits. Variants A, B, and C use different soft weights but exactly the same hard predicates.",
                "Exact-credit selection is a bounded dynamic programme over the remaining credit target. If C is the target and n is the number of candidates, its pseudo-polynomial cost is O(nC) before diversity checks. Prerequisite closure and cycle detection are O(|V|+|E|). Candidate retrieval dominates only when embeddings are not cached; with a vector index, approximate retrieval is sublinear in repository size and exact reranking is applied to the bounded shortlist.",
                "The final guarantee is conditional and explicit. If the verifier accepts a persisted plan, then the implemented predicates for credit volume, semester load, forward prerequisites, level, duplicate titles, domain quotas, regulatory closure, and minimum real-course LO evidence all hold for that stored version. This is a software-verifiable structural guarantee, not a claim that accreditation or future student attainment has been proven."
            ])
        ],
        "5. Experimental Design": [
            ("5.1. Why Several Metrics Are Required", [
                "Classification and ranking answer different questions. ROC-AUC measures pair separation across thresholds, whereas PR-AUC is more informative when supported links are sparse. F1 evaluates one frozen decision threshold. Recall@K asks whether expert-relevant courses appear in the shortlist shown to a programme designer. MRR rewards placing the first relevant course early, and nDCG gives larger gain to strongly rated items near the top.",
                "Plan-level validity cannot be inferred from pair metrics. A model may retrieve relevant courses yet generate an impossible credit total or reverse a prerequisite. The evaluation therefore follows a hierarchy: pair classification, shortlist ranking, graph and semester consistency, complete-plan invariants, and qualitative expert review. A result is promoted only if the relevant lower-level metrics improve and the higher-level invariants remain intact.",
                "The reported GOSO comparison is an explanatory cohort analysis rather than a causal estimate. The cohorts have equal size and use the same verifier, but programme content is not randomly assigned to regulatory mode. Confidence intervals and blinded multi-institutional review are planned for a larger prospective evaluation."
            ])
        ],
    },
    "ru": {
        "1. Введение": [
            ("1.1. Что такое ЕПВО и почему эти данные важны", [
                "ЕПВО — принятое сокращение государственной Единой платформы высшего образования Республики Казахстан. Реестр образовательных программ внутри ЕПВО является государственной информационной системой, куда вузы загружают программы для экспертизы и регистрации. Примерно за десять лет здесь сформирована уникальная продольная база уже разработанных программ и внешних экспертных мнений. Это не обычный каталог: корпус содержит уровень образования, направление и группу ОП, цели, РО, дисциплины, кредиты, рекомендуемые семестры, описания и заявленные связи дисциплин с результатами обучения.",
                "Доказательные данные возникают в академическом жизненном цикле. Сначала кафедра формулирует цель программы и результаты обучения, собирает учебный план и указывает, какие дисциплины обеспечивают каждый результат. Затем внешние эксперты соответствующего направления изучают программу и оценивают достижимость заявленных связей как низкую, среднюю или высокую. В восстановленном массиве этим уровням соответствуют числа 0, 0,5 и 1. Голоса сохраняются и агрегируются без превращения средней и высокой связи в одинаковую бинарную метку.",
                "Пример поясняет различие сигналов. Для пары «Цифровая криминалистика — сбор и сохранение цифровых доказательств» SBERT оценивает семантическое соответствие двух текстов. Оценка ЕПВО отвечает на другой вопрос: насколько сильно эксперты поддерживали сходные связи в ранее проверенных программах. Curriculum-KAG показывает оба значения отдельно, а дисциплина допускается только после проверки уровня, направления, кредитов, пререквизитов и профессионального контекста.",
                "Такое происхождение делает корпус ценнее непроверенного списка названий, однако не превращает каждое историческое решение в универсальную истину. Программы разрабатывались разными вузами, в разные годы и для разных условий. Поэтому ЕПВО используется как экспертная память и доказательство, а не как оракул. Риск механического запоминания снижают разбиение по целым программам, консервативный перенос оценок, независимый verifier и подтверждение специалистом.",
                "ГОСО РК — государственный общеобязательный стандарт образования Республики Казахстан. Он задаёт защищённые компоненты для конкретного уровня образования. В режиме соответствия Казахстану они обязательны, хотя не всегда являются обычными профильными дисциплинами ЕПВО. Поэтому международный и нормативный режимы сравниваются как по полному плану, так и по профильной части после корректного исключения защищённого блока."
            ])
        ],
        "2. Связанные исследования": [
            ("2.1. Почему система называется KAG", [
                "Retrieval-Augmented Generation (RAG) обычно находит релевантные фрагменты документов и передаёт их генеративной модели как контекст. Это снижает число вымышленных фактов, но сами фрагменты не задают, какая дисциплина является пререквизитом, какой РО она обеспечивает, какой эксперт подтвердил связь и соблюдены ли кредиты всего плана. RAG хорошо отвечает по найденным документам, но не гарантирует корректную образовательную программу.",
                "Knowledge-Augmented Generation (KAG) дополняет поиск структурированным знанием. В Curriculum-KAG слой знаний содержит канонические дисциплины, уровни и направления программ, градуированные экспертные связи дисциплина-РО, пре- и постреквизиты, предметные области, нормативный статус, кредиты, типичные семестры и происхождение данных. Кандидаты расширяются и фильтруются через граф, после чего constrained-планировщик собирает программу, а verifier проверяет сохранённый результат целиком.",
                "Поэтому KAG лучше обычного RAG именно для задачи проектирования учебного плана, хотя не заявляется универсальным победителем во всех языковых задачах. Его преимущество — работа с отношениями и правилами: система объясняет выбор дисциплины, показывает, что изучается до и после неё, отличает реальную карточку ЕПВО от синтетического bridge и отклоняет красивый текстовый ответ, нарушающий академические ограничения. Настоящая работа развивает исходную концепцию Curriculum-KAG [20] до универсального генератора классических и междисциплинарных программ на стыке двух областей."
            ])
        ],
        "3. Данные и происхождение": [
            ("3.1. Формальная модель данных", [
                "Программа p задаётся уровнем образования h_p, направлением d_p, группой образовательных программ b_p, текстом цели, множеством результатов O_p и исходным набором дисциплин C_p. Raw-запись дисциплины содержит название, многоязычное описание, кредиты, тип компонента, типичный семестр и заявленные связи с РО. Нормализованная дисциплина является канонической сущностью, связанной с одной или несколькими исходными карточками. Эта связь многие-к-одному сохраняет происхождение и не допускает многократного включения дублей в план.",
                "Экспертное наблюдение представлено кортежем (p,c,l,e), где дисциплина c оценивалась относительно исходного результата l программы p, а e принадлежит множеству {0; 0,5; 1}. Одна каноническая дисциплина может получить разные оценки в разных программах. Метод не превращает контекстное распределение в единственную глобальную истину: сначала используется память в рамках программы или близкого направления, а общий агрегат служит резервным сигналом.",
                "Рабочий репозиторий создаётся последовательно: raw EPVO -> normalized EPVO -> утверждённый многоязычный репозиторий -> кандидаты выбранного проекта. Исходные объекты и контрольные суммы неизменяемы. Нормализация приводит пробелы, коды, кредиты и многоязычные fingerprints к единому виду. Утверждение создаёт пригодную каноническую карточку, а project scope фильтрует уровень, направление и группу до начала семантического ранжирования.",
                "Dataset Passport фиксирует время извлечения, checksum источника, версию схемы, полноту языков, дубли, распределение оценок, split и ограничения. Model Benchmark сохраняет идентификатор модели, seed, порог, checksum конфигурации и точные идентификаторы программ в train, validation и test."
            ]),
            ("3.2. Восстановление экспертных оценок и контроль утечки", [
                "При первоначальной нормализации существовал риск сворачивания внешней оценки в бинарную связь. Исправленный конвейер сохраняет градуированные значения. Всего 876 842 исходных голоса агрегированы в 800 793 контекстные пары дисциплина-РО; они являются признаком доказательности и целью эксперимента, но не автоматическим решением о включении дисциплины.",
                "Карточки одной образовательной программы никогда не разделяются между обучением и тестом. Program-disjoint split строже случайного разбиения пар, поскольку дисциплины и РО одной программы имеют общий словарь и структуру. Иначе почти дублированный контекст дал бы завышенную оценку обобщения. Пороги, веса смеси и early stopping определяются по validation-программам и замораживаются перед тестом.",
                "Отрицательные примеры по возможности выбираются внутри того же уровня образования и близкого направления. Они сложнее случайных междоменных пар и лучше соответствуют реальной задаче ранжирования: нужно различить несколько правдоподобных дисциплин, а не отделить медицину от несвязанного инженерного РО. Экспериментальные programme-scoped memory и ranking-loss ветви оцениваются отдельно от рабочей модели."
            ])
        ],
        "4.4. Формальная целевая функция и независимый verifier": [
            ("4.5. Переменные оптимизации, гарантии и сложность", [
                "Пусть x_i — бинарная переменная выбора дисциплины, а z_is — её назначение в семестр s. Допустимое множество фиксируется до оптимизации проверками уровня образования, направления, группы, контекста и минимального доказательства связи с РО. Следовательно, планировщик работает только с академически правдоподобными кандидатами и не может компенсировать недопустимую дисциплину большим весом целевой функции.",
                "Целевая функция сочетает ограниченное покрытие РО, семантическую и экспертную релевантность, баланс областей, разнообразие, согласование семестра и минимизацию bridge. Жёсткие ограничения не входят в взвешенную сумму. Рост similarity не может «купить» нарушение пререквизита, неверный уровень образования или недобор кредитов. Варианты A, B и C используют разные мягкие веса, но одинаковые обязательные предикаты.",
                "Точный подбор кредитов реализован динамическим программированием по остаточному объёму. При цели C и n кандидатах псевдополиномиальная сложность равна O(nC) до проверки разнообразия. Замыкание пререквизитов и поиск циклов имеют сложность O(|V|+|E|). Поиск в репозитории доминирует только при отсутствии кэша embeddings; векторный индекс формирует короткий список, после чего точный reranking выполняется на ограниченном числе кандидатов.",
                "Финальная гарантия условна и проверяема. Если verifier принимает сохранённый план, то для этой версии выполнены реализованные предикаты кредитов, нагрузки семестров, направленных вперёд пререквизитов, уровня, отсутствия дублей, квот областей, нормативного замыкания и минимального покрытия РО реальными дисциплинами. Это структурная программная гарантия, но не утверждение о полученной аккредитации или будущей успеваемости студентов."
            ])
        ],
        "5. Дизайн эксперимента": [
            ("5.1. Почему требуется несколько метрик", [
                "Классификация и ranking отвечают на разные вопросы. ROC-AUC измеряет разделение пар по всем порогам, а PR-AUC информативнее при редких подтверждённых связях. F1 оценивает один замороженный порог. Recall@K показывает, попали ли экспертно релевантные дисциплины в короткий список разработчика. MRR поощряет раннюю позицию первой правильной дисциплины, а nDCG учитывает весь градуированный порядок.",
                "Качество полного плана нельзя вывести из метрик пар. Модель может найти релевантные дисциплины, но собрать неверную сумму кредитов или развернуть пререквизит назад. Поэтому оценивание образует иерархию: классификация пар, ranking кандидатов, граф и семестры, инварианты полного плана и качественная экспертиза. Результат внедряется только при улучшении нужных нижних метрик и сохранении верхнеуровневых ограничений.",
                "Сравнение ГОСО является объясняющим когортным анализом, а не причинной оценкой. Когорты равны по размеру и используют один verifier, однако содержание программ не назначалось случайно. Для следующей расширенной проверки запланированы доверительные интервалы и слепая межвузовская экспертиза."
            ])
        ],
    },
}


DATA_ROWS_EN = [
    ["Educational programmes", "12,215"],
    ["Raw course records", "408,638"],
    ["Raw learning outcomes", "124,521"],
    ["Normalized courses", "191,292"],
    ["Normalized expert links", "932,483"],
    ["Pairs with numerical expert labels", "800,793"],
    ["Approved operational courses", "21,525"],
]
DATA_ROWS_RU = [
    ["Образовательные программы", "12 215"],
    ["Исходные записи дисциплин", "408 638"],
    ["Исходные результаты обучения", "124 521"],
    ["Нормализованные дисциплины", "191 292"],
    ["Нормализованные экспертные связи", "932 483"],
    ["Пары с числовой экспертной оценкой", "800 793"],
    ["Утверждённые рабочие дисциплины", "21 525"],
]

MODEL_ROWS = [
    ["Base multilingual SBERT", "0.7031", "0.6946", "0.6886", "-"],
    ["Fine-tuned SBERT 40k", "0.7666", "0.7681", "0.7188", "operational"],
    ["Cached candidate 12k", "0.7748", "0.7746", "0.7268", "experimental"],
]

REFERENCES = [
    "[1] Biggs, J. (1996). Enhancing teaching through constructive alignment. Higher Education, 32, 347-364.",
    "[2] Harden, R. M. (2007). Outcome-based education: the future is today. Medical Teacher, 29(7), 625-629.",
    "[3] Reimers, N., & Gurevych, I. (2019). Sentence-BERT: Sentence embeddings using Siamese BERT-networks. EMNLP-IJCNLP, 3982-3992.",
    "[4] Robertson, S., & Zaragoza, H. (2009). The probabilistic relevance framework: BM25 and beyond. Foundations and Trends in Information Retrieval, 3(4), 333-389.",
    "[5] Lewis, P., et al. (2020). Retrieval-augmented generation for knowledge-intensive NLP tasks. Advances in Neural Information Processing Systems, 33, 9459-9474.",
    "[6] Deb, K., Pratap, A., Agarwal, S., & Meyarivan, T. (2002). A fast and elitist multiobjective genetic algorithm: NSGA-II. IEEE Transactions on Evolutionary Computation, 6(2), 182-197.",
    "[7] Kipf, T. N., & Welling, M. (2017). Semi-supervised classification with graph convolutional networks. ICLR.",
    "[8] Hamilton, W. L., Ying, R., & Leskovec, J. (2017). Inductive representation learning on large graphs. NeurIPS, 1024-1034.",
    "[9] Hochreiter, S., & Schmidhuber, J. (1997). Long short-term memory. Neural Computation, 9(8), 1735-1780.",
    "[10] Gebru, T., et al. (2021). Datasheets for datasets. Communications of the ACM, 64(12), 86-92.",
    "[11] Mitchell, M., et al. (2019). Model cards for model reporting. FAT* 2019, 220-229.",
    "[12] Unified Higher Education Platform of the Republic of Kazakhstan. Registry of Educational Programs. https://epvo.kz/#/register/education_program",
    "[13] Ministry of Science and Higher Education of the Republic of Kazakhstan. State Compulsory Standards of Higher and Postgraduate Education, revision dated May 4, 2026.",
    "[14] ABET. (2025). Criteria for Accrediting Computing Programs.",
    "[15] Crawley, E. F., Malmqvist, J., Östlund, S., Brodeur, D. R., & Edström, K. (2014). Rethinking Engineering Education: The CDIO Approach. Springer.",
    "[16] González, J., & Wagenaar, R. (Eds.). (2003). Tuning Educational Structures in Europe. University of Deusto.",
    "[17] Jantassova, D., Kozhanov, M., & Shebalina, O. (2021). Digital platform as a tool for internationalization: Model for formation of international competences' database applying the hierarchy analysis method. Journal of Theoretical and Applied Information Technology, 99(21), 4942-4957. https://www.scopus.com/pages/publications/85119289073",
    "[18] Kozhanov, M., Mosavi, A., Amirov, A., Kaibassova, D., Poser, V., Shakhatova, A., Lira, L., & Eigner, G. (2024). Syllabus design and planning with artificial intelligence and the potential of generative AI. In 2024 IEEE 24th International Symposium on Computational Intelligence and Informatics, 257-264. https://doi.org/10.1109/CINTI63048.2024.10830766",
    "[19] Kozhanov, M., & Várkonyi-Kóczy, A. R. (2025). Modelling curricula with GNN and LSTM for link and sequence prediction. Eurasian Journal of Mathematical and Computer Applications, 13(4), 159-167. https://doi.org/10.32523/2306-6172-2025-13-4-159-167; https://www.scopus.com/pages/publications/105027578723",
    "[20] Kozhanov, M., Makó, C., Póser, V., Eigner, G., & Mosavi, A. (2026). Knowledge Augmented Generation for Curriculum Planning. Eurasian Journal of Mathematical and Computer Applications, 14(1), 17-34. https://doi.org/10.32523/2306-6172-2026-14-1-17-34; https://www.scopus.com/pages/publications/105035219019",
    "[21] Najjar-Ghabel, S., Daneshpour, A., Amirov, A., Yousefimehr, B., & Kozhanov, M. (2026). Deep learning and gradient boosting for predicting educational attainment based on socioeconomic and regional features. Eurasian Journal of Mathematical and Computer Applications, 14(2), 95-114. https://doi.org/10.32523/2306-6172-2026-14-2-95-114; https://www.scopus.com/pages/publications/105043909947",
]


FIGURES = {
    1: (
        "Why Curriculum-KAG extends plain RAG.",
        "Create a publication-grade two-panel vector comparison for a Q2 scientific journal, landscape 16:9, white background. Left panel titled 'RAG': a user query enters a vector search, three unstructured text chunks are retrieved, and an LLM produces a plausible answer; show dashed warning icons for missing prerequisite relations, credit validation, expert provenance, and regulatory rules. Right panel titled 'Curriculum-KAG': programme goal and learning outcomes enter hybrid SBERT+BM25 retrieval; retrieved courses connect to a typed knowledge graph with nodes for courses, LOs, domains, semesters, experts, and regulations; arrows pass through constrained planning and an independent verifier to one or three valid curriculum variants. Add a bottom strip: 'retrieval + relations + provenance + constraints + verification'. Use navy for retrieval, teal for knowledge, amber for constraints, green for verified output. No robots, brains, neon, 3D, or invented metrics. All labels in English; editable SVG plus 300-dpi PNG; readable at one journal-column width.",
        "Почему Curriculum-KAG расширяет обычный RAG.",
        "Создай двухпанельную векторную схему уровня Q2, горизонталь 16:9, белый фон. Слева панель «RAG»: запрос поступает в векторный поиск, возвращаются три неструктурированных фрагмента, LLM формирует правдоподобный ответ; пунктирными предупреждениями показать отсутствие пререквизитных связей, проверки кредитов, происхождения экспертной оценки и нормативных правил. Справа «Curriculum-KAG»: цель и РО поступают в гибридный SBERT+BM25; дисциплины связываются с типизированным графом узлов «дисциплина, РО, область, семестр, эксперт, норматив»; далее constrained planner, независимый verifier и один или три валидных плана. Внизу формула-смысл: «поиск + отношения + происхождение + ограничения + проверка». Цвета: navy для поиска, teal для знаний, amber для ограничений, green для принятого результата. Без роботов, мозга, неона, 3D и вымышленных чисел. Русские подписи, SVG и PNG 300 dpi, читаемость в одну колонку.",
    ),
    2: (
        "Algorithm 1: EPVO evidence formation, normalization, and leakage-safe split.",
        "Create a precise scientific flowchart for Algorithm 1, landscape 16:9, white background. Use six numbered stages. (1) University department: goal, programme LOs, curriculum, declared course-to-LO links. (2) External field experts: independently assign low=0, medium=0.5, high=1 to each declared relation. (3) Immutable raw EPVO: original JSON cards and SHA-256 checksum. (4) Normalization: language cleaning, canonical_title, credit and level codes, multilingual fingerprint, deduplication while preserving all source edges. (5) Aggregation: retain the graded mean by programme-course-source-LO; explicitly show that 0/0.5/1 are not binarised. (6) Grouped split: entire programmes go to Train, Validation, or Test, with no programme crossing partitions. Finish with Approved RU/KK/EN Repository and Dataset Passport. Show measured scale cards: 12,215 programmes, 408,638 raw courses, 124,521 LOs, 932,483 normalized links, 800,793 graded pairs. Separate 'expert evidence' from 'AI prediction'. Flat vector, navy/teal/green, no decorative AI imagery, editable SVG and 300-dpi PNG.",
        "Алгоритм 1: формирование доказательств ЕПВО, нормализация и split без утечки.",
        "Создай точную научную блок-схему алгоритма 1, горизонталь 16:9, белый фон, шесть нумерованных стадий. (1) Кафедра: цель, РО, учебный план и заявленные связи дисциплина-РО. (2) Внешние эксперты направления независимо ставят низкая=0, средняя=0,5, высокая=1. (3) Неизменяемый raw EPVO: исходные JSON и SHA-256. (4) Нормализация: языки, canonical_title, кредиты, коды уровней, многоязычный fingerprint, дедупликация с сохранением всех исходных рёбер. (5) Агрегация по программе-дисциплине-исходному РО без бинаризации 0/0,5/1. (6) Grouped split: целые программы уходят в Train, Validation или Test и не пересекаются. Финал: утверждённый RU/KK/EN-репозиторий и Dataset Passport. Вставить карточки масштаба: 12 215 программ, 408 638 raw-дисциплин, 124 521 РО, 932 483 нормализованные связи, 800 793 градуированные пары. Визуально разделить «экспертное доказательство» и «прогноз ИИ». Плоский вектор, navy/teal/green, SVG и PNG 300 dpi.",
    ),
    3: (
        "Algorithm 2: evidence-aware course-to-learning-outcome scoring.",
        "Create a left-to-right scientific pipeline for Algorithm 2. Inputs: one new programme LO and only courses from the selected education level, EPVO direction, and programme group. Stage A: multilingual text preparation for RU/KK/EN title and description. Stage B: SBERT embeddings and cosine similarity in parallel with normalized BM25 exact-term retrieval. Stage C: calibrated AI probability. Stage D: retrieve historical expert votes 0/0.5/1 only from sufficiently similar source LOs; show conservative transfer threshold. Stage E: hard guards for level 6B/7M/8D, domain scope, foreign professional context, duplicates, and minimum LO evidence. Stage F: rank candidates and display separate AI score, EPVO expert score, final relationship strength, source programme, and expert-feedback buttons. Include a red rejected-course path and a green admitted-course path. Emphasize that AI% and EPVO% are independent evidence channels and are not added. White background, flat vector, concise English labels, navy/teal/amber/green/red, no invented numbers, editable SVG and 300-dpi PNG.",
        "Алгоритм 2: доказательное оценивание связи дисциплина-РО.",
        "Создай слева направо научный pipeline алгоритма 2. Вход: один РО новой программы и только дисциплины выбранного уровня, направления и группы ЕПВО. A: подготовка RU/KK/EN названия и описания. B: параллельно SBERT embeddings с cosine similarity и нормированный BM25 для точных терминов. C: калиброванная вероятность ИИ. D: исторические экспертные оценки 0/0,5/1 переносятся только через достаточно похожие исходные РО; показать порог консервативного переноса. E: жёсткие guards уровня 6B/7M/8D, scope области, чужого профессионального контекста, дублей и минимального LO evidence. F: ranking с отдельными AI score, EPVO score, итоговой силой, исходной программой и кнопками экспертной обратной связи. Красная ветвь отклонённой дисциплины и зелёная принятой. Подчеркнуть, что AI% и EPVO% не складываются. Белый фон, плоский вектор, русские подписи, navy/teal/amber/green/red, без вымышленных чисел, SVG и PNG 300 dpi.",
    ),
    4: (
        "Algorithm 3: exact-credit, real-course-first repair and bridge control.",
        "Create a publication-grade decision flow for Algorithm 3. Start with total programme credits and programme mode. If Kazakhstan requirements are enabled, insert the protected level-specific block for bachelor's, master's, or doctorate and subtract its credits. Feed the remaining professional-credit target into a dynamic-programming matrix whose columns are reachable credit sums and rows are admissible real EPVO courses with immutable credits and evidence utility. Highlight the highest-evidence exact combination. Then perform prerequisite closure and bounded whole-course swaps. Only if no exact real-course solution exists, open a clearly marked bridge branch: create an interdisciplinary bridge for an uncovered LO, show three real replacement candidates, require expert confirmation, and never call a generic elective placeholder a course. End with exact total credits, minimum bridge count, and verifier pass/fail. Use a compact formula inset for D_i(k), but keep it readable. White background, navy for real courses, amber for repair, purple for bridge, green for verified result, red for failure; English labels; editable SVG and 300-dpi PNG.",
        "Алгоритм 3: точные кредиты, приоритет реальных дисциплин и bridge-контроль.",
        "Создай научную decision-flow схему алгоритма 3. Начало: общий объём кредитов и режим программы. Если включены требования Казахстана, вставить защищённый блок конкретного уровня — бакалавриат, магистратура или докторантура — и вычесть его кредиты. Остаток профильных кредитов подать в матрицу динамического программирования: столбцы — достижимые суммы, строки — допустимые реальные дисциплины ЕПВО с неизменяемыми кредитами и evidence utility. Выделить точную комбинацию с максимальным доказательством, затем выполнить замыкание пререквизитов и ограниченные обмены целых дисциплин. Только при отсутствии точного реального решения открыть явно помеченную bridge-ветвь: междисциплинарный модуль для непокрытого РО, три реальные замены и обязательное подтверждение эксперта; общий «компонент по выбору» не считать дисциплиной. Финал: точные кредиты, минимум bridge и verifier pass/fail. Вставить компактную формулу D_i(k). Белый фон, navy/amber/purple/green/red, русские подписи, SVG и PNG 300 dpi.",
    ),
    5: (
        "Algorithm 4: graph-aware semester scheduling and transactional A/B/C generation.",
        "Create a detailed semester-planning diagram for Algorithm 4, landscape 16:9. Left: a directed acyclic graph with foundation, intermediate, advanced, practice, research, and final-attestation course nodes; prerequisite arrows always point forward. Centre: calculate earliest feasible semester from topological depth, typical EPVO semester, semantic floor/ceiling, and a 27-33 credit load band. Show bounded whole-course moves and swaps; forbid credit rewriting and moving an advanced course into semester 1. Right: generate either one requested plan or variants A/B/C with different soft ranking weights but identical hard constraints. Below, an independent verifier recomputes exact credits, semester loads, level, domain quotas, LO evidence, duplicates, regulatory closure, and prerequisite order. Atomic commit occurs only if every requested variant passes and A/B/C fingerprints differ; otherwise the previous active plan remains. Use English labels, white background, navy/teal/amber/green/red, no decorative AI, editable SVG and 300-dpi PNG.",
        "Алгоритм 4: графовое распределение по семестрам и транзакционная генерация A/B/C.",
        "Создай подробную схему semester planning алгоритма 4, горизонталь 16:9. Слева направленный ацикличный граф узлов «основа, промежуточный, продвинутый курс, практика, исследование, итоговая аттестация»; стрелки пререквизитов только вперёд. В центре рассчитать ранний допустимый семестр по топологической глубине, типовому семестру ЕПВО, semantic floor/ceiling и диапазону 27–33 кредита. Показать ограниченные переносы и обмены целых дисциплин; запретить изменение кредитов и перенос продвинутого курса в первый семестр. Справа генерируется один выбранный план либо A/B/C с разными мягкими ranking-весами и одинаковыми жёсткими правилами. Внизу независимый verifier пересчитывает кредиты, нагрузку, уровень, квоты, LO evidence, дубли, нормативное замыкание и порядок пререквизитов. Atomic commit только если всё прошло и fingerprints A/B/C различаются; иначе старый активный план сохраняется. Русские подписи, белый фон, navy/teal/amber/green/red, SVG и PNG 300 dpi.",
    ),
    6: (
        "International mode and Kazakhstan regulatory requirements.",
        "Create a Q2-journal comparison chart with two aligned curriculum lanes and a small bar panel. Lane A titled 'International mode': no Kazakhstan protected block, semantic-adjusted semester alignment 86.84%, profile provenance 100%, hard violations 0. Lane B titled 'Kazakhstan requirements (SCES RK/GOSO)': protected mandatory components, practices, research work, and final attestation; 35 protected elements across three reported plans; semantic-adjusted alignment 81.23%; profile provenance after regulatory exclusion 100%; hard violations 0. Show the shared professional-course layer in navy and the protected local layer in amber. Display forward prerequisite arrows and 27-33 credit semester bands. Add a textual note: the 5.61-point gap represents reduced scheduling freedom, not lower quality of profile courses. Do not call either mode accredited. White background, restrained colours, editable SVG and 300-dpi PNG, English labels.",
        "Международный режим и местные требования Казахстана.",
        "Создай сравнительный график уровня Q2 с двумя дорожками учебного плана и небольшой панелью столбцов. A «Международный режим»: без защищённого блока Казахстана, semantic-adjusted semester alignment 86,84%, профильное происхождение 100%, жёсткие нарушения 0. B «Требования Казахстана (ГОСО РК)»: обязательные компоненты, практики, НИР и итоговая аттестация; 35 защищённых элементов в трёх отчётных планах; alignment 81,23%; профильное происхождение после исключения нормативного блока 100%; нарушения 0. Общий профильный слой — navy, местный защищённый — amber. Показать стрелки пререквизитов и диапазон 27–33 кредита. Добавить пояснение: разница 5,61 п.п. означает меньшую свободу расписания, а не худшее качество профильных дисциплин. Не называть режимы аккредитованными. Белый фон, SVG и PNG 300 dpi, русские подписи.",
    ),
    7: (
        "Curriculum-KAG programme-design interface: required screenshots.",
        "Create a four-panel composite from REAL APPLICATION SCREENSHOTS, not generated UI. Panel A: the programme wizard showing classical/interdisciplinary selector, bachelor's/master's/doctorate level, one or two EPVO directions/groups, programme language, credits, goal, and the AI buttons for three goals and learning outcomes. Panel B: a generated semester plan with plan A/B/C selector, credit totals, cycle/component labels, optional course descriptions, and expandable quality sections. Panel C: the 'Why selected?' course card showing full LO text, AI score, EPVO expert score, final strength, source, cycle, prerequisites/postrequisites, and expert confirmation buttons. Panel D: the futuristic semester graph with courses on the left, prerequisite/postrequisite arrows in the centre, and semester competencies on the right. Crop personal data, keep browser chrome out, use one consistent interface language, 300-dpi PNG; add thin navy borders and panel labels A-D without altering screen content.",
        "Интерфейс проектирования Curriculum-KAG: необходимые скриншоты.",
        "Собери четырёхпанельную композицию из РЕАЛЬНЫХ СКРИНШОТОВ приложения, не генерируй интерфейс. A: мастер программы с выбором классическая/междисциплинарная, бакалавриат/магистратура/докторантура, одно или два направления и группы ЕПВО, язык, кредиты, цель и отдельные AI-кнопки предложения трёх целей и РО. B: сформированный план по семестрам с переключателем A/B/C, суммой кредитов, циклом/компонентом, галочкой описаний и раскрывающимися проверками качества. C: карточка «Почему выбрана?» с полным РО, AI score, EPVO score, итоговой силой, источником, циклом, пре- и постреквизитами и экспертными кнопками. D: футуристичный граф — дисциплины слева, стрелки пре- и постреквизитов в центре, компетенции семестра справа. Обрезать персональные данные и рамку браузера, использовать один язык, PNG 300 dpi, тонкие navy-рамки и метки A-D без изменения содержимого.",
    ),
    8: (
        "Registry comparison and quality-verification interface: required screenshots.",
        "Create a two-panel composite from REAL APPLICATION SCREENSHOTS. Left panel: 'Compare with EPVO' for one generated interdisciplinary programme, showing several similar registry programmes, common and missing courses, Jaccard/containment indicators, typical courses, unique elements, and an action for adding a priority course. Right panel: the plan-quality dashboard with exact credits, semester load, prerequisite violations, LO coverage sources, domain compliance, international checklist, Kazakhstan requirements when enabled, Dataset Passport, and Model Benchmark. Include a visible plan A/B/C selector and explanatory text that checklist completion is not accreditation. Use the same programme and language in both panels, hide usernames and identifiers, crop empty margins, do not invent values, output 300-dpi PNG suitable for a two-column journal page.",
        "Сравнение с Реестром и интерфейс проверки качества: необходимые скриншоты.",
        "Собери двухпанельную композицию из РЕАЛЬНЫХ СКРИНШОТОВ. Слева экран «Сравнить с ЕПВО» для одной междисциплинарной программы: несколько похожих программ Реестра, общие и отсутствующие дисциплины, Jaccard/containment, типовые курсы, уникальные элементы и действие добавления приоритетной дисциплины. Справа dashboard качества: точные кредиты, нагрузка семестров, нарушения пререквизитов, источники покрытия РО, соответствие областям, международный чек-лист, требования Казахстана при включённом режиме, Dataset Passport и Model Benchmark. Обязательно показать переключатель A/B/C и пояснение, что прохождение чек-листа не является аккредитацией. В обеих панелях одна программа и один язык; скрыть имя пользователя и идентификаторы, убрать пустые поля, не выдумывать значения, PNG 300 dpi для страницы двухколоночного журнала.",
    ),
}


EQUATIONS = {
    "en": {
        "3. Data and Provenance": [
            ("y_(pcl)=1/|E_(pcl)| ∑_(e∈E_(pcl))e", "Here E_pcl is the set of external votes for programme p, course c, and source outcome l; y_pcl retains the graded mean on the interval [0,1]."),
            ("P_a∩P_b=∅, ∀a≠b", "P_train, P_val, and P_test are disjoint sets of complete educational programmes, not randomly mixed course-LO rows."),
        ],
        "4. Method": [
            ("u_i=g_θ(x_i)/‖g_θ(x_i)‖_2", "The multilingual encoder g_theta maps the text x_i of course i to the L2-normalised vector u_i; outcomes are encoded analogously as v_j."),
            ("s_(ij)^sem=u_i^T v_j", "The semantic score is cosine similarity because both embeddings have unit length."),
            ("r_(ij)=αs_(ij)^sem+(1-α)s_(ij)^lex", "The retrieval score combines dense semantic evidence and a normalised BM25 lexical score; alpha is selected on validation programmes."),
            ("p_(ij)=σ((r_(ij)-τ)/T)", "Validation-selected threshold tau and temperature T convert the retrieval score into a calibrated AI confidence p_ij."),
        ],
        "4.1. Auditable Admission and Domain Evidence": [
            ("a_(jl)=|K(o_j)∩K(l)|/|K(o_j)∪K(l)|", "K(.) extracts key concepts; a_jl prevents an expert score from being transferred between outcomes that only share a course title."),
            ("e_(ij)=max_l{y_(pil)a_(jl):a_(jl)≥τ_e}", "Transferred EPVO evidence e_ij is the strongest contextually admissible product of historical expert strength and outcome similarity."),
            ("A_i=I_i^level I_i^scope I_i^LO I_i^context", "The Boolean admission indicator A_i equals one only when level, programme scope, LO evidence, and professional-context guards all pass."),
            ("w_(ij)=max(p_(ij),e_(ij),f_(ij))", "Final link strength w_ij uses the strongest admissible channel: calibrated AI prediction, transferred EPVO evidence, or confirmed local expert feedback f_ij."),
        ],
        "4.2. Exact Credits, Bridge Control, and Regulatory Closure": [
            ("C^*-δ≤∑_i c_i x_i≤C^*+δ", "C* is the programme credit target, delta is the declared tolerance, c_i is the immutable credit value of a real course, and x_i indicates selection."),
            ("D_i(k)=max{D_(i-1)(k),D_(i-1)(k-c_i)+q_i}", "The exact-fit recurrence retains the highest evidence value q_i for every reachable professional-credit remainder k."),
            ("B(P)=∑_i x_i I_i^bridge", "B(P) counts visible synthetic bridge units; the planner minimises this quantity after attempting exact-credit combinations of real EPVO courses."),
        ],
        "4.3. Knowledge Graph, Semester Repair, and A/B/C Transactions": [
            ("G=(V_C∪V_O,E_pre∪E_LO∪E_sim)", "The knowledge graph joins course and outcome nodes with prerequisite, coverage, and semantic-similarity edges."),
            ("∑_s z_(is)=x_i", "Every selected course is assigned to exactly one semester; an unselected course is assigned to none."),
            ("L_s≤∑_i c_i z_(is)≤U_s", "The total load in semester s must remain between its lower and upper limits."),
            ("z_(ia)z_(jb)=1⇒a<b", "Every prerequisite course i must precede its dependent course j; concurrent placement is not accepted for ordinary course edges."),
            ("s_i^*=arg min_(s∈F_i)|s-t_i|", "The selected semester is the nearest feasible member of F_i to the EPVO typical semester t_i; F_i already enforces load, prerequisite, and pedagogical bounds."),
        ],
        "4.4. Formal Objective and Independent Verifier": [
            ("C_j(P)=1-∏_(i:x_i=1)(1-w_(ij))", "Probabilistic coverage C_j(P) aggregates independent supporting courses while remaining bounded by one."),
            ("F_cov(P)=1/m ∑_(j=1)^m min{C_j(P)/τ_cov,1}", "Bounded mean coverage stops an already satisfied outcome from compensating for an uncovered one."),
            ("Φ(P)=λ_1F_cov+λ_2Rel+λ_3Div+λ_4Sem", "The positive utility combines bounded coverage, relevance, diversity, and semester agreement."),
            ("P^*=arg max_(P∈F){Φ(P)-λ_5Red-λ_6B(P)}", "Optimisation is restricted to the hard-feasible set F; redundancy and synthetic bridges are penalised."),
            ("V(P)=∏_(k=1)^K I_k(P)", "The independent verifier accepts a persisted plan only when every hard predicate I_k is true."),
        ],
        "5. Experimental Design": [
            ("Recall@K=|Rel(q)∩TopK(q)|/|Rel(q)|", "Recall@K is the share of all expert-relevant courses retrieved among the first K candidates for query q."),
            ("MRR=1/|Q| ∑_(q∈Q)1/rank_q", "MRR is the mean reciprocal position of the first relevant course; earlier correct results receive greater weight."),
            ("DCG@K=∑_(r=1)^K(2^(g_r)-1)/log_2(r+1)", "DCG uses graded relevance g_r, so a highly supported course at an early rank contributes more."),
            ("nDCG@K=DCG@K/IDCG@K", "nDCG divides the observed gain by the ideal ordering and therefore lies between zero and one."),
            ("D(A,B)=1-|A∩B|/|A∪B|", "Pairwise Jaccard distance D(A,B) checks that curriculum alternatives are not identical lists with different labels."),
        ],
    },
    "ru": {
        "3. Данные и происхождение": [
            ("y_(pcl)=1/|E_(pcl)| ∑_(e∈E_(pcl))e", "E_pcl — множество внешних голосов для программы p, дисциплины c и исходного результата l; y_pcl сохраняет градуированное среднее на интервале [0;1]."),
            ("P_a∩P_b=∅, ∀a≠b", "P_train, P_val и P_test являются непересекающимися множествами целых образовательных программ, а не случайно смешанных строк дисциплина-РО."),
        ],
        "4. Метод": [
            ("u_i=g_θ(x_i)/‖g_θ(x_i)‖_2", "Многоязычный encoder g_theta преобразует текст x_i дисциплины i в L2-нормированный вектор u_i; РО аналогично кодируется вектором v_j."),
            ("s_(ij)^sem=u_i^T v_j", "Семантическая оценка равна косинусному сходству, поскольку оба embedding имеют единичную длину."),
            ("r_(ij)=αs_(ij)^sem+(1-α)s_(ij)^lex", "Retrieval-score сочетает плотное семантическое доказательство и нормированный лексический BM25; alpha выбирается на validation-программах."),
            ("p_(ij)=σ((r_(ij)-τ)/T)", "Порог tau и температура T, выбранные на validation, преобразуют retrieval-score в калиброванную уверенность ИИ p_ij."),
        ],
        "4.1. Аудируемый допуск и доменные доказательства": [
            ("a_(jl)=|K(o_j)∩K(l)|/|K(o_j)∪K(l)|", "K(.) извлекает ключевые понятия; a_jl не позволяет переносить экспертную оценку между РО, совпавшими только по названию дисциплины."),
            ("e_(ij)=max_l{y_(pil)a_(jl):a_(jl)≥τ_e}", "Перенесённое доказательство ЕПВО e_ij — максимальное допустимое произведение исторической силы связи и сходства формулировок РО."),
            ("A_i=I_i^level I_i^scope I_i^LO I_i^context", "Булев индикатор допуска A_i равен единице только при прохождении уровня, scope программы, доказательства РО и профессионального контекста."),
            ("w_(ij)=max(p_(ij),e_(ij),f_(ij))", "Итоговая сила w_ij использует наиболее сильный допустимый канал: прогноз ИИ, перенесённое доказательство ЕПВО или подтверждённую локальную обратную связь f_ij."),
        ],
        "4.2. Точные кредиты, bridge-контроль и нормативное замыкание": [
            ("C^*-δ≤∑_i c_i x_i≤C^*+δ", "C* — требуемый объём программы, delta — заданный допуск, c_i — неизменяемые кредиты реальной дисциплины, x_i — признак выбора."),
            ("D_i(k)=max{D_(i-1)(k),D_(i-1)(k-c_i)+q_i}", "Рекуррентность exact-fit сохраняет комбинацию с максимальной доказательностью q_i для каждого достижимого остатка профессиональных кредитов k."),
            ("B(P)=∑_i x_i I_i^bridge", "B(P) считает явно обозначенные синтетические bridge-модули; система минимизирует их после поиска точной комбинации реальных дисциплин ЕПВО."),
        ],
        "4.3. Граф знаний, semester-repair и транзакция A/B/C": [
            ("G=(V_C∪V_O,E_pre∪E_LO∪E_sim)", "Граф знаний объединяет узлы дисциплин и РО с рёбрами пререквизитов, покрытия и семантического сходства."),
            ("∑_s z_(is)=x_i", "Каждая выбранная дисциплина назначается ровно в один семестр, а невыбранная не назначается."),
            ("L_s≤∑_i c_i z_(is)≤U_s", "Суммарная нагрузка семестра s остаётся между нижней и верхней границами."),
            ("z_(ia)z_(jb)=1⇒a<b", "Пререквизит i обязан предшествовать зависимой дисциплине j; одновременное размещение обычного ребра не принимается."),
            ("s_i^*=arg min_(s∈F_i)|s-t_i|", "Выбирается ближайший к типовому семестру ЕПВО t_i допустимый семестр из F_i, где уже учтены нагрузка, пререквизиты и педагогические границы."),
        ],
        "4.4. Формальная целевая функция и независимый verifier": [
            ("C_j(P)=1-∏_(i:x_i=1)(1-w_(ij))", "Вероятностное покрытие C_j(P) агрегирует несколько подтверждающих дисциплин и не превышает единицу."),
            ("F_cov(P)=1/m ∑_(j=1)^m min{C_j(P)/τ_cov,1}", "Ограниченное среднее не позволяет уже закрытому РО компенсировать полностью непокрытый результат."),
            ("Φ(P)=λ_1F_cov+λ_2Rel+λ_3Div+λ_4Sem", "Положительная полезность объединяет покрытие, релевантность, разнообразие и согласование семестров."),
            ("P^*=arg max_(P∈F){Φ(P)-λ_5Red-λ_6B(P)}", "Оптимизация выполняется только на жёстко допустимом множестве F; дублирование и синтетические bridge штрафуются."),
            ("V(P)=∏_(k=1)^K I_k(P)", "Независимый verifier принимает сохранённый план только при истинности каждого обязательного предиката I_k."),
        ],
        "5. Дизайн эксперимента": [
            ("Recall@K=|Rel(q)∩TopK(q)|/|Rel(q)|", "Recall@K — доля всех экспертно релевантных дисциплин, попавших в первые K кандидатов для запроса q."),
            ("MRR=1/|Q| ∑_(q∈Q)1/rank_q", "MRR — средняя обратная позиция первой релевантной дисциплины; ранний правильный ответ получает больший вес."),
            ("DCG@K=∑_(r=1)^K(2^(g_r)-1)/log_2(r+1)", "DCG использует градуированную релевантность g_r, поэтому сильная связь на ранней позиции вносит больший вклад."),
            ("nDCG@K=DCG@K/IDCG@K", "nDCG делит наблюдаемый выигрыш на идеальный порядок и находится от нуля до единицы."),
            ("D(A,B)=1-|A∩B|/|A∪B|", "Попарное расстояние Жаккара D(A,B) проверяет, что варианты планов не являются одинаковыми списками с разными буквами."),
        ],
    },
}


ALGORITHMS = {
    "en": {
        "3. Data and Provenance": ("EPVO normalization and leakage-safe split", "Immutable raw EPVO records and source checksums.", [
            "Validate schema and preserve each source JSON object without mutation.",
            "Normalize titles, language variants, credits, levels, directions, groups, and component codes.",
            "Construct canonical-course fingerprints; retain every source-record edge and expert vote.",
            "Aggregate votes by programme, canonical course, and source LO without binarising 0/0.5/1.",
            "Group all entities by programme identifier, then assign complete programmes to train, validation, or test.",
            "Approve multilingual canonical records and calculate a dataset-passport checksum.",
        ], "Normalized entities, graded expert evidence, disjoint splits, and an auditable approved repository."),
        "4. Method": ("Evidence-aware course-to-LO scoring", "New programme outcomes and project-scoped candidate courses.", [
            "Encode course and LO texts with the frozen multilingual SBERT encoder.",
            "Retrieve a bounded shortlist using dense similarity and normalized BM25.",
            "Transfer EPVO expert evidence only through sufficiently similar source outcomes.",
            "Apply education-level, direction/group, professional-context, and duplicate guards.",
            "Combine admissible AI, EPVO, and confirmed local-expert evidence without adding their percentages.",
            "Return ranked candidates with provenance and an explanation for every score.",
        ], "A ranked and auditable course-to-LO candidate set."),
        "4.2. Exact Credits, Bridge Control, and Regulatory Closure": ("Exact-credit real-course-first repair", "Protected regulatory core, professional credit target, and admissible courses.", [
            "Subtract protected SCES RK components from the total target when regulatory mode is enabled.",
            "Run dynamic programming over immutable real-course credit values and evidence utilities.",
            "Select the highest-evidence exact remainder; close prerequisite dependencies.",
            "If no exact real-course combination exists, search bounded whole-course swaps.",
            "Create a labelled bridge only when a professional LO or remaining credit gap is still infeasible.",
            "Reject the candidate if the final credit tolerance or LO-evidence invariant fails.",
        ], "An exact-credit course set with the minimum defensible number of bridges."),
        "4.3. Knowledge Graph, Semester Repair, and A/B/C Transactions": ("Transactional graph-aware curriculum synthesis", "Selected courses, prerequisite graph, semester limits, typical semesters, and variant weights.", [
            "Topologically order the prerequisite graph and calculate the earliest feasible semester for each course.",
            "Apply semantic floors and ceilings for foundation, advanced, practice, research, and final-attestation units.",
            "Assign each course to the closest feasible typical semester while maintaining the credit band.",
            "Repair hard violations first; then perform bounded swaps that improve semester agreement.",
            "Generate the requested one or three variants and verify each from persisted-form data.",
            "Commit all variants atomically only if every invariant passes and fingerprints are distinct.",
        ], "One valid plan or a distinct, fully verified A/B/C set; otherwise no stored plan is replaced."),
    },
    "ru": {
        "3. Данные и происхождение": ("Нормализация ЕПВО и split без утечки", "Неизменяемые raw-записи ЕПВО и checksum источника.", [
            "Проверить схему и сохранить каждый исходный JSON без изменения.",
            "Нормализовать названия, языки, кредиты, уровни, направления, группы и коды компонентов.",
            "Построить fingerprints канонических дисциплин; сохранить рёбра ко всем источникам и экспертные голоса.",
            "Агрегировать оценки по программе, дисциплине и исходному РО без бинаризации 0/0,5/1.",
            "Сгруппировать сущности по идентификатору программы и целиком назначить программы в train, validation или test.",
            "Утвердить многоязычные карточки и вычислить checksum Dataset Passport.",
        ], "Нормализованные сущности, градуированные оценки, непересекающиеся split и аудируемый репозиторий."),
        "4. Метод": ("Доказательное оценивание связи дисциплина-РО", "РО новой программы и кандидаты из project-scoped репозитория.", [
            "Закодировать тексты дисциплин и РО замороженным многоязычным SBERT.",
            "Получить ограниченный shortlist по dense similarity и нормированному BM25.",
            "Перенести оценку ЕПВО только через достаточно сходные исходные РО.",
            "Применить фильтры уровня, направления/группы, профессионального контекста и дублей.",
            "Объединить допустимые AI, EPVO и подтверждённые локальным экспертом сигналы без сложения процентов.",
            "Вернуть ranking с происхождением и объяснением каждой оценки.",
        ], "Ранжированный и аудируемый набор кандидатов дисциплина-РО."),
        "4.2. Точные кредиты, bridge-контроль и нормативное замыкание": ("Точная сумма: сначала реальные дисциплины", "Защищённое нормативное ядро, цель профильных кредитов и допустимые курсы.", [
            "В режиме ГОСО вычесть защищённые компоненты из общего целевого объёма.",
            "Запустить динамическое программирование по неизменяемым кредитам и доказательности реальных дисциплин.",
            "Выбрать наиболее сильный точный остаток и выполнить замыкание пререквизитов.",
            "Если точной комбинации нет, проверить ограниченные обмены целых дисциплин.",
            "Создать явно помеченный bridge только при неустранимом пробеле профессионального РО или кредитов.",
            "Отклонить вариант при нарушении допуска кредитов или инварианта покрытия РО.",
        ], "Точная сумма кредитов с минимальным обоснованным числом bridge."),
        "4.3. Граф знаний, semester-repair и транзакция A/B/C": ("Транзакционный графовый синтез плана", "Дисциплины, граф пререквизитов, пределы нагрузки, типовые семестры и веса варианта.", [
            "Топологически упорядочить граф и вычислить ранний допустимый семестр каждой дисциплины.",
            "Применить семантические границы для основ, продвинутых курсов, практик, исследований и аттестации.",
            "Назначить ближайший допустимый типовой семестр с сохранением кредитного диапазона.",
            "Сначала исправить жёсткие нарушения, затем выполнить ограниченные обмены для улучшения согласования.",
            "Сформировать один или три запрошенных варианта и независимо проверить каждый.",
            "Атомарно зафиксировать варианты только при всех инвариантах и различных fingerprints.",
        ], "Один валидный план или различный проверенный набор A/B/C; при ошибке старый план не заменяется."),
    },
}


def build(template: Path, output: Path, metrics: dict, lang: str):
    shutil.copy2(template, output)
    document = Document(output)
    remove_after_front_matter(document)
    data = EN if lang == "en" else RU
    set_front(
        document,
        data["title"],
        "Murat Kozhanov",
        [
            "Academy of Public Administration under the President of the Republic of Kazakhstan, Astana, Kazakhstan",
            "Doctoral School of Applied Informatics and Applied Mathematics, Óbuda University, Budapest, Hungary",
        ],
        "Corresponding author: m.kozhanov@ktu.edu.kz",
    )
    add_abstract(
        document,
        data["abstract"],
        data["keywords"],
        ("Abstract", "Keywords") if lang == "en" else ("Аннотация", "Ключевые слова"),
    )

    equation_number = 1

    def insert_figure(number: int):
        cap, prompt = (
            (FIGURES[number][0], FIGURES[number][1])
            if lang == "en"
            else (FIGURES[number][2], FIGURES[number][3])
        )
        add_figure_prompt(document, number, cap, prompt, lang)

    for heading, paragraphs in data["sections"]:
        add_heading(document, heading)
        for text in paragraphs:
            add_body(document, text)
        for subheading, inserted_paragraphs in SCIENTIFIC_INSERTIONS[lang].get(heading, []):
            add_heading(document, subheading, level=2)
            for text in inserted_paragraphs:
                add_body(document, text)
        for equation, explanation in EQUATIONS[lang].get(heading, []):
            add_equation(document, equation, equation_number)
            add_equation_explanation(document, explanation)
            equation_number += 1
        if heading in ("2. Related Work", "2. Связанные исследования"):
            insert_figure(1)
        if heading in ("3. Data and Provenance", "3. Данные и происхождение"):
            rows = DATA_ROWS_EN if lang == "en" else DATA_ROWS_RU
            add_table(
                document,
                ["Dataset entity", "Count"] if lang == "en" else ["Сущность набора данных", "Количество"],
                rows,
                "Table 1. EPVO corpus and operational repository." if lang == "en" else "Таблица 1. Корпус ЕПВО и рабочий репозиторий.",
            )
            insert_figure(2)
        if heading in ("4. Method", "4. Метод"):
            insert_figure(3)
        if heading in (
            "4.2. Exact Credits, Bridge Control, and Regulatory Closure",
            "4.2. Точные кредиты, bridge-контроль и нормативное замыкание",
        ):
            insert_figure(4)
        if heading in (
            "4.3. Knowledge Graph, Semester Repair, and A/B/C Transactions",
            "4.3. Граф знаний, semester-repair и транзакция A/B/C",
        ):
            insert_figure(5)
        if heading in ("6. Results", "6. Результаты"):
            add_table(
                document,
                ["Model", "ROC-AUC", "PR-AUC", "F1", "Status"] if lang == "en" else ["Модель", "ROC-AUC", "PR-AUC", "F1", "Статус"],
                MODEL_ROWS,
                "Table 2. Course-to-learning-outcome classification." if lang == "en" else "Таблица 2. Классификация связей дисциплина-результат обучения.",
            )
            headers, rows, caption = cohort_table(metrics, lang)
            add_table(document, headers, rows, caption)
            insert_figure(6)
        if heading in (
            "8. System Implementation and Reproducibility",
            "8. Реализация и воспроизводимость",
        ):
            insert_figure(7)
            insert_figure(8)

    add_heading(document, "Acknowledgements" if lang == "en" else "Благодарности")
    add_body(
        document,
        "This research and the scientific internship at Óbuda University were supported by a grant from the private Shakhmardan Yessenov Science and Education Foundation, founded by Galimzhan Yessenov, under the 2026 Scientific Internships in World Laboratories programme. The author also thanks the academic experts who provided qualitative feedback during system demonstrations. No external accreditation decision is claimed."
        if lang == "en"
        else "Исследование и научная стажировка в Óbuda University выполнены при поддержке гранта Частного фонда «Научно-образовательный фонд Shakhmardan Yessenov Foundation», учредителем которого является Галимжан Есенов, в рамках программы «Научные стажировки в лабораториях мира — 2026». Автор также благодарит академических экспертов, предоставивших качественную обратную связь во время демонстраций системы. Формальное аккредитационное решение в работе не заявляется.",
    )
    add_heading(document, "References" if lang == "en" else "Литература")
    for reference in REFERENCES:
        paragraph = document.add_paragraph(style="Els-body-text")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(0)
        add_formatted_text(paragraph, reference, size=9)

    # Ensure Word refreshes fields if the journal template adds them later.
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    document.core_properties.title = data["title"]
    document.core_properties.author = "Murat Kozhanov"
    document.core_properties.subject = "Curriculum-KAG"
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--metrics", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    metrics = json.loads(args.metrics.read_text(encoding="utf-8"))
    args.output_dir.mkdir(parents=True, exist_ok=True)
    build(args.template, args.output_dir / "Curriculum_KAG_TEM_Journal_EN.docx", metrics, "en")
    build(args.template, args.output_dir / "Curriculum_KAG_TEM_Journal_RU.docx", metrics, "ru")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
