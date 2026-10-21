"""Strict, provenance-preserving parsing for university baseline files."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import zipfile
from typing import Any


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_XLSX_MEMBERS = 1000
REQUIRED_FIELDS = ("code", "title", "credits")


class ImportValidationError(ValueError):
    """Raised when an import violates the explicit baseline contract."""


def _fail(message: str) -> None:
    raise ImportValidationError(message)


def _validate_columns(headers: list[Any], columns: dict[str, str]) -> dict[str, int]:
    if tuple(columns.keys()) != REQUIRED_FIELDS or len(columns) != len(REQUIRED_FIELDS):
        _fail("columns must explicitly map code, title, and credits")
    normalized = [str(header).strip() if header is not None else "" for header in headers]
    if len(normalized) != len(set(normalized)):
        _fail("duplicate source columns are not allowed")
    positions: dict[str, int] = {}
    for field in REQUIRED_FIELDS:
        source_name = columns[field]
        if not isinstance(source_name, str) or not source_name.strip():
            _fail(f"missing mapping for {field}")
        try:
            positions[field] = normalized.index(source_name.strip())
        except ValueError:
            _fail(f"missing source column for {field}")
    return positions


def _parse_row(values: list[Any], positions: dict[str, int], source_row: int) -> dict[str, Any]:
    def value(field: str) -> Any:
        item = values[positions[field]] if positions[field] < len(values) else None
        if isinstance(item, str) and item.startswith("="):
            _fail(f"formula is not allowed at source row {source_row}")
        return item

    code = str(value("code") or "").strip()
    title = str(value("title") or "").strip()
    if not code or not title:
        _fail(f"empty code or title at source row {source_row}")
    raw_credits = value("credits")
    try:
        credits = float(raw_credits)
    except (TypeError, ValueError):
        _fail(f"invalid credits at source row {source_row}")
    if not math.isfinite(credits) or credits <= 0:
        _fail(f"invalid credits at source row {source_row}")
    return {
        "source_row": source_row,
        "original_code": code,
        "original_title": title,
        "credits": credits,
    }


def _parse_csv(data: bytes, columns: dict[str, str]) -> list[dict[str, Any]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        _fail(f"CSV must be UTF-8: {exc}")
    reader = csv.reader(io.StringIO(text))
    try:
        headers = next(reader)
    except StopIteration:
        _fail("source is empty")
    positions = _validate_columns(headers, columns)
    rows: list[dict[str, Any]] = []
    for source_row, values in enumerate(reader, start=2):
        if not any(str(item).strip() for item in values):
            continue
        if len(rows) >= MAX_ROWS:
            _fail(f"more than {MAX_ROWS} data rows are not allowed")
        rows.append(_parse_row(values, positions, source_row))
    if not rows:
        _fail("source contains no data rows")
    return rows


def _check_xlsx_limits(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                _fail(f"XLSX archive has more than {MAX_XLSX_MEMBERS} members")
            uncompressed = sum(member.file_size for member in members)
            if uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                _fail("XLSX archive exceeds uncompressed size limit")
    except zipfile.BadZipFile:
        _fail("invalid XLSX archive")


def _parse_xlsx(data: bytes, columns: dict[str, str]) -> list[dict[str, Any]]:
    _check_xlsx_limits(data)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        _fail(f"XLSX support unavailable: {exc}")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=False)
        header_cells = next(iterator)
    except (StopIteration, ValueError, OSError, zipfile.BadZipFile) as exc:
        _fail(f"invalid XLSX source: {exc}")
    headers = [cell.value for cell in header_cells]
    positions = _validate_columns(headers, columns)
    rows: list[dict[str, Any]] = []
    for source_row, cells in enumerate(iterator, start=2):
        values: list[Any] = []
        for cell in cells:
            if cell.data_type == "f":
                _fail(f"formula is not allowed at source row {source_row}")
            values.append(cell.value)
        if not any(str(item).strip() for item in values if item is not None):
            continue
        if len(rows) >= MAX_ROWS:
            _fail(f"more than {MAX_ROWS} data rows are not allowed")
        rows.append(_parse_row(values, positions, source_row))
    workbook.close()
    if not rows:
        _fail("source contains no data rows")
    return rows


def parse_source(data: bytes, filename: str, columns: dict[str, str]) -> list[dict[str, Any]]:
    """Parse a UTF-8 CSV or XLSX using only the caller's explicit column mapping."""
    if not isinstance(data, (bytes, bytearray)):
        _fail("source data must be bytes")
    if len(data) > MAX_UPLOAD_BYTES:
        _fail(f"source exceeds {MAX_UPLOAD_BYTES} byte upload limit")
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        return _parse_csv(bytes(data), columns)
    if suffix == "xlsx":
        return _parse_xlsx(bytes(data), columns)
    _fail("only CSV and XLSX sources are supported")


def canonical_hash(rows: list[dict[str, Any]]) -> str:
    payload = json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def validate_mapping(
    rows: list[dict[str, Any]],
    mappings: dict[int, int],
    catalogue: dict[int, dict[str, Any]],
) -> list[dict[str, Any]]:
    """Resolve every imported row to an existing catalogue course without fuzzy matching."""
    source_rows = {int(row["source_row"]) for row in rows}
    if set(mappings) != source_rows:
        _fail("mapping must contain exactly one entry for every source row")
    resolved: list[dict[str, Any]] = []
    for row in rows:
        source_row = int(row["source_row"])
        course_id = mappings[source_row]
        if not isinstance(course_id, int) or isinstance(course_id, bool) or course_id not in catalogue:
            _fail(f"mapping references unknown catalogue course at source row {source_row}")
        course = catalogue[course_id]
        content = {key: course[key] for key in sorted(course) if key not in {"id", "course_id"}}
        content_sha256 = hashlib.sha256(
            json.dumps(content, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")
        ).hexdigest()
        resolved.append({**row, "course_id": course_id, "content_sha256": content_sha256})
    return resolved
