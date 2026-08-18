from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy.exc import SQLAlchemyError
from app.database import get_db
from app.models.user import User
from app.services.auth import get_current_user
from app.models.plan import Plan, PlanItem
from app.models.project import ProjectVersion
from app.models.course import Course
from app.models.bridge_module import BridgeModule
from app.models.embedding import MatchScore
from app.models.audit import AuditEvent
from app.services.content_localization import course_localization_map
from app.services.language import normalize_language
from app.kag.knowledge_graph import get_graph_stats
from app.kag.embedding_service import embedding_service
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO

router = APIRouter()


def localized_title(localizations: dict, course: Course, language: str = "ru") -> str:
    translations = (localizations.get(course.id) or {}).get("title_translations") or {}
    language = normalize_language(language)
    # Prefer the requested interface language, then use a deterministic fallback
    # so an incomplete legacy row never produces an empty export cell.
    return (
        translations.get(language)
        or translations.get("ru")
        or translations.get("kk")
        or translations.get("en")
        or course.title
    )


@router.post("/{project_version_id}")
async def export_plan(
    project_version_id: int,
    format: str = "xlsx",
    language: str = "ru",
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """Export curriculum plan to XLSX"""
    
    if format != "xlsx":
        raise HTTPException(status_code=400, detail="Сейчас поддерживается только формат XLSX")
    
    try:
        # Get project version
        project_version = db.query(ProjectVersion).filter(
            ProjectVersion.id == project_version_id
        ).first()
        
        if not project_version:
            raise HTTPException(status_code=404, detail="Версия проекта не найдена")
        
        # Get best plan (variant A)
        plan = db.query(Plan).filter(
            Plan.project_version_id == project_version_id,
            Plan.variant_type == "A"
        ).first()
        
        if not plan:
            raise HTTPException(status_code=404, detail="Учебный план не найден")
        
        # Create workbook
        wb = openpyxl.Workbook()
        
        # Sheet 1: Curriculum Plan
        ws1 = wb.active
        ws1.title = "Curriculum Plan"
        
        # Headers
        headers = ["Semester", "Course Code", "Course Title", "Credits", "Domain", "Type", "Prerequisites"]
        ws1.append(headers)
        
        # Style headers
        for cell in ws1[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
        
        # Get plan items
        items = db.query(PlanItem).filter(PlanItem.plan_id == plan.id).order_by(PlanItem.semester).all()
        course_ids = [item.course_id for item in items if item.course_id]
        localizations = course_localization_map(db, course_ids)
        
        for item in items:
            if item.course_id:
                course = db.query(Course).filter(Course.id == item.course_id).first()
                ws1.append([
                    item.semester,
                    course.course_id,
                    localized_title(localizations, course, language),
                    course.credits,
                    course.domain,
                    item.course_type,
                    ", ".join([p.course_id for p in course.prerequisites]) if course.prerequisites else ""
                ])
            elif item.bridge_module_id:
                bm = db.query(BridgeModule).filter(BridgeModule.id == item.bridge_module_id).first()
                ws1.append([
                    item.semester,
                    bm.course_id,
                    bm.title + " [BRIDGE]",
                    bm.credits,
                    "Interdisciplinary",
                    item.course_type,
                    ", ".join(bm.prerequisites) if bm.prerequisites else ""
                ])
        
        # Sheet 2: LO Coverage Matrix
        ws2 = wb.create_sheet("LO Coverage")
        
        # Get LOs
        los = project_version.learning_outcomes
        ws2.append(["Course Code", "Course Title"] + [lo.lo_code for lo in los])
        
        # Get all courses in plan
        courses = db.query(Course).filter(Course.id.in_(course_ids)).all()
        
        for course in courses:
            row = [course.course_id, localized_title(localizations, course, language)]
            for lo in los:
                match = db.query(MatchScore).filter(
                    MatchScore.course_id == course.id,
                    MatchScore.lo_id == lo.id
                ).first()
                row.append(f"{match.score:.2f}" if match else "0.00")
            ws2.append(row)
        
        # Sheet 3: Metrics
        ws3 = wb.create_sheet("Metrics")
        ws3.append(["Metric", "Value"])
        
        metrics = plan.metrics_json
        for key, value in metrics.items():
            ws3.append([key.replace("_", " ").title(), value])
        
        # Sheet 4: Deterministic verification details
        ws4 = wb.create_sheet("Verification")
        ws4.append(["Check", "Result"])
        verification = (metrics or {}).get("verification", {})
        for key in ["feasible", "hard_violation_count", "target_credits", "total_credits", "maximum_total_credits", "min_lo_coverage", "average_lo_coverage", "coverage_threshold", "evidence_count", "redundancy"]:
            ws4.append([key, str(verification.get(key, ""))])
        ws4.append(["prerequisite_violations", len(verification.get("prerequisite_violations", []))])
        ws4.append(["semester_load_violations", len(verification.get("semester_load_violations", []))])

        ws5 = wb.create_sheet("Knowledge Graph")
        ws5.append(["Metric", "Value"])
        for key, value in {**get_graph_stats(db), **embedding_service.get_status()}.items():
            ws5.append([key, str(value)])

        ws6 = wb.create_sheet("Audit Log")
        ws6.append(["Timestamp", "Action", "Entity Type", "Entity ID", "Details"])
        for event in db.query(AuditEvent).order_by(AuditEvent.id.desc()).limit(200).all():
            ws6.append([str(event.timestamp), event.action, event.entity_type, event.entity_id, str(event.details_json or {})])

        ws7 = wb.create_sheet("International Quality")
        ws7.append(["Framework Score", str((metrics or {}).get("international_quality", {}).get("score", ""))])
        ws7.append(["Passed", str((metrics or {}).get("international_quality", {}).get("passed", ""))])
        ws7.append([])
        ws7.append(["Check", "Passed", "Evidence", "Recommendation"])
        for check in (metrics or {}).get("international_quality", {}).get("checks", []):
            ws7.append([
                check.get("name", ""),
                str(check.get("passed", "")),
                check.get("evidence", ""),
                check.get("recommendation", ""),
            ])
        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename=curriculum_plan_{project_version_id}.xlsx"
            }
        )
        
    except (SQLAlchemyError, OSError, ValueError, TypeError) as e:
        raise HTTPException(status_code=500, detail=f"Не удалось выполнить экспорт: {e.__class__.__name__}") from e

