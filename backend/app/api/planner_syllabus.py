from __future__ import annotations

from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.audit import AuditEvent
from app.models.bridge_module import BridgeModule
from app.models.course import Course, CourseChunk
from app.models.embedding import MatchScore
from app.models.syllabus import SyllabusDraft
from app.models.user import User
from app.services.auth import get_current_user
from app.services.content_localization import (
    course_localization_map,
    course_localization_payload,
)


router = APIRouter()

@router.get("/syllabus/{kind}/{entity_id}")
async def generate_course_syllabus(
    kind: str,
    entity_id: int,
    weeks: int = 15,
    contact_share: float = 0.5,
    mode: str = "academic",
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Generate an auditable syllabus draft whose hours equal credits × 30."""
    if weeks < 10 or weeks > 20:
        raise HTTPException(status_code=400, detail="Количество недель должно быть от 10 до 20")
    if contact_share < 0.2 or contact_share > 0.8:
        raise HTTPException(status_code=400, detail="Доля контактной работы должна быть от 0,2 до 0,8")
    if mode not in {"academic", "practical", "project"}:
        raise HTTPException(status_code=400, detail="Режим должен быть академическим, практическим или проектным")

    if kind == "course":
        entity = db.query(Course).filter(Course.id == entity_id).first()
        if not entity:
            raise HTTPException(status_code=404, detail="Дисциплина не найдена")
        localization = course_localization_payload(db, entity.id)
        title_translations = localization.get("title_translations", {})
        description_translations = localization.get("description_translations", {})
        code = entity.course_id
        title = title_translations.get("ru") or title_translations.get("kk") or title_translations.get("en") or entity.title
        description = description_translations.get("ru") or description_translations.get("kk") or description_translations.get("en") or entity.description
        credits, topics = int(entity.credits or 0), list(entity.topics or [])
        outcomes, assessments = list(entity.learning_outcomes or []), list(entity.assessment_methods or [])
        related_localizations = course_localization_map(db, [item.id for item in list(entity.prerequisites) + list(entity.postrequisites)])
        prerequisites = [{
            "code": item.course_id,
            "title": item.title,
            "title_translations": related_localizations.get(item.id, {}).get("title_translations", {}),
        } for item in entity.prerequisites]
        postrequisites = [{
            "code": item.course_id,
            "title": item.title,
            "title_translations": related_localizations.get(item.id, {}).get("title_translations", {}),
        } for item in entity.postrequisites]
    elif kind == "bridge":
        entity = db.query(BridgeModule).filter(BridgeModule.id == entity_id).first()
        if not entity:
            raise HTTPException(status_code=404, detail="Bridge-модуль не найден")
        code, title, description = entity.course_id, entity.title, entity.description
        credits, topics = int(entity.credits or 0), list(entity.topics or [])
        outcomes, assessments = list(entity.learning_outcomes or []), list(entity.assessment_methods or [])
        prerequisites, postrequisites = [], []
    else:
        raise HTTPException(status_code=400, detail="Тип должен быть дисциплиной или bridge-модулем")

    if credits <= 0:
        raise HTTPException(status_code=400, detail="Количество кредитов дисциплины должно быть положительным")
    if not topics:
        topics = [f"{title}: foundational concepts", f"{title}: methods and tools", f"{title}: applied project"]
    if not outcomes:
        outcomes = [f"Explain core concepts of {title}", f"Apply methods of {title} to a practical task"]
    if not assessments:
        assessments = ["Practical assignments", "Midterm assessment", "Final project or examination"]

    total_hours = credits * 30
    contact_hours = round(total_hours * contact_share)
    independent_hours = total_hours - contact_hours
    lecture_share = {"academic": 0.55, "practical": 0.35, "project": 0.20}[mode]
    lecture_hours = round(contact_hours * lecture_share)
    practical_hours = contact_hours - lecture_hours

    def distribute(total, count):
        base, remainder = divmod(int(total), int(count))
        return [base + (1 if index < remainder else 0) for index in range(count)]

    lectures = distribute(lecture_hours, weeks)
    practicals = distribute(practical_hours, weeks)
    independent = distribute(independent_hours, weeks)
    rows = []
    midpoint = max(2, (weeks + 1) // 2)
    mode_activity = {
        "academic": "Guided analysis and seminar discussion",
        "practical": "Laboratory or applied case exercise",
        "project": "Project workshop and supervised implementation",
    }
    for index in range(weeks):
        week = index + 1
        topic_index = min(len(topics) - 1, int(index * len(topics) / weeks))
        topic = str(topics[topic_index])
        if week == 1:
            phase, activity, evidence = "foundation", "Diagnostic task and guided introduction", "Diagnostic response"
        elif week == midpoint:
            phase, activity, evidence = "midterm", "Integrated midterm task", "Assessed midterm artefact"
        elif week == weeks:
            phase, activity, evidence = "final", "Final synthesis and defence", "Final project or examination evidence"
        elif week >= weeks - 2:
            phase, activity, evidence = "integration", mode_activity[mode], "Integrated draft or project increment"
        else:
            phase, activity = "development", mode_activity[mode]
            evidence = "Laboratory report" if mode == "practical" else "Project increment" if mode == "project" else "Analytical assignment"
        if week == midpoint:
            assessment = assessments[min(1, len(assessments) - 1)]
        elif week == weeks:
            assessment = assessments[-1]
        elif week % 3 == 0:
            assessment = assessments[0]
        else:
            assessment = "Formative feedback"
        rows.append({
            "week": week,
            "phase": phase,
            "topic": topic,
            "lecture_hours": lectures[index],
            "practical_hours": practicals[index],
            "independent_hours": independent[index],
            "total_hours": lectures[index] + practicals[index] + independent[index],
            "learning_outcome": str(outcomes[index % len(outcomes)]),
            "learning_activity": activity,
            "assessment": str(assessment),
            "evidence": evidence,
        })

    mapped_outcomes = {row["learning_outcome"] for row in rows}
    validations = {
        "hours_match": sum(row["total_hours"] for row in rows) == total_hours,
        "all_outcomes_mapped": all(str(outcome) in mapped_outcomes for outcome in outcomes),
        "midterm_present": any(row["phase"] == "midterm" for row in rows),
        "final_assessment_present": any(row["phase"] == "final" for row in rows),
        "evidence_present_every_week": all(bool(row["evidence"]) for row in rows),
    }
    validations["ready_for_expert_review"] = all(validations.values())

    return {
        "kind": kind, "entity_id": entity_id, "code": code, "title": title,
        "title_translations": title_translations if kind == "course" else {},
        "description": description,
        "description_translations": description_translations if kind == "course" else {},
        "credits": credits, "weeks": weeks, "mode": mode,
        "hours_per_credit": 30, "total_hours": total_hours,
        "contact_share": contact_share, "contact_hours": contact_hours,
        "lecture_hours": lecture_hours, "practical_hours": practical_hours,
        "independent_hours": independent_hours, "learning_outcomes": outcomes,
        "assessment_methods": assessments, "prerequisites": prerequisites,
        "postrequisites": postrequisites, "thematic_plan": rows, "validations": validations,
        "assumptions": [
            "One academic credit equals 30 academic hours of total student workload.",
            "The contact/independent split is a configurable institutional template.",
            "The draft requires academic-methodological review before approval.",
        ],
        "hours_check": validations["hours_match"],
    }


@router.get("/syllabus/draft/{kind}/{entity_id}")
async def get_syllabus_draft(kind: str, entity_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    draft = db.query(SyllabusDraft).filter(SyllabusDraft.kind == kind, SyllabusDraft.entity_id == entity_id, SyllabusDraft.created_by == current_user.id).order_by(SyllabusDraft.id.desc()).first()
    if not draft:
        raise HTTPException(status_code=404, detail="Сохранённый черновик силлабуса не найден")
    return {"id": draft.id, "status": draft.status, "updated_at": draft.updated_at, "content": draft.content_json}


@router.post("/syllabus/draft/{kind}/{entity_id}")
async def save_syllabus_draft(kind: str, entity_id: int, payload: dict, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    content = payload.get("content") if isinstance(payload, dict) else None
    rows = content.get("thematic_plan", []) if isinstance(content, dict) else []
    if not rows:
        raise HTTPException(status_code=400, detail="Тематический план пуст")
    expected = int(content.get("credits", 0)) * int(content.get("hours_per_credit", 30))
    actual = sum(int(row.get("total_hours", 0)) for row in rows)
    if expected <= 0 or actual != expected:
        raise HTTPException(status_code=400, detail=f"Недельная нагрузка {actual} ч. не соответствует требуемой {expected} ч.")
    draft = db.query(SyllabusDraft).filter(SyllabusDraft.kind == kind, SyllabusDraft.entity_id == entity_id, SyllabusDraft.created_by == current_user.id).order_by(SyllabusDraft.id.desc()).first()
    if not draft:
        draft = SyllabusDraft(kind=kind, entity_id=entity_id, created_by=current_user.id)
        db.add(draft)
    draft.weeks = int(content.get("weeks", len(rows)))
    draft.contact_share = float(content.get("contact_share", 0.5))
    draft.mode = str(content.get("mode", "academic"))
    draft.content_json = content
    draft.status = "draft"
    db.add(AuditEvent(user_id=current_user.id, action="save_syllabus_draft", entity_type=kind, entity_id=entity_id, details_json={"weeks": draft.weeks, "mode": draft.mode, "hours": actual}))
    db.commit(); db.refresh(draft)
    return {"id": draft.id, "status": draft.status, "updated_at": draft.updated_at}


@router.post("/syllabus/export-docx")
async def export_syllabus_docx(payload: dict, current_user: User = Depends(get_current_user)):
    from docx import Document
    content = payload.get("content") if isinstance(payload, dict) else None
    if not isinstance(content, dict) or not content.get("thematic_plan"):
        raise HTTPException(status_code=400, detail="Содержание силлабуса пусто")
    document = Document()
    document.add_heading("Рабочая программа дисциплины (Syllabus)", 0)
    document.add_heading(str(content.get("title", "Дисциплина")), level=1)
    document.add_paragraph(f"Код: {content.get('code', '')} | Кредиты: {content.get('credits', '')} | Недель: {content.get('weeks', '')} | Всего часов: {content.get('total_hours', '')}")
    document.add_heading("Описание", level=2); document.add_paragraph(str(content.get("description") or "Не заполнено"))
    document.add_heading("Результаты обучения", level=2)
    for outcome in content.get("learning_outcomes", []): document.add_paragraph(str(outcome), style="List Number")
    document.add_heading("Тематический план", level=2)
    table = document.add_table(rows=1, cols=7); table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, ["Неделя", "Тема", "Деятельность", "Часы", "LO", "Оценивание", "Доказательство"]): cell.text = title
    for row in content["thematic_plan"]:
        values = [row.get("week"), row.get("topic"), row.get("learning_activity"), row.get("total_hours"), row.get("learning_outcome"), row.get("assessment"), row.get("evidence")]
        for cell, value in zip(table.add_row().cells, values): cell.text = str(value or "")
    document.add_paragraph("Черновик требует проверки и утверждения преподавателем или методистом.")
    output = BytesIO(); document.save(output); output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=course_syllabus.docx"})


@router.get("/{plan_id}/evidence-bundle")
async def get_evidence_bundle(
    plan_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Return the exact LO-to-course evidence used by a persisted plan."""
    import hashlib
    import json
    from datetime import datetime, timezone
    from app.models.plan import Plan

    plan = db.query(Plan).filter(Plan.id == plan_id).first()
    if not plan:
        raise HTTPException(status_code=404, detail="Учебный план не найден")
    course_ids = sorted({item.course_id for item in plan.items if item.course_id is not None})
    version = plan.project_version
    evidence = []
    for lo in version.learning_outcomes:
        rows = db.query(MatchScore).filter(
            MatchScore.project_version_id == version.id,
            MatchScore.lo_id == lo.id,
            MatchScore.course_id.in_(course_ids or [-1]),
        ).order_by(MatchScore.score.desc(), MatchScore.course_id).all()
        evidence.append({
            "lo_code": lo.lo_code,
            "lo_text": lo.lo_text,
            "matches": [{
                "course_id": row.course_id,
                "chunk_id": row.chunk_id,
                "chunk_text": db.query(CourseChunk.chunk_text).filter(CourseChunk.id == row.chunk_id).scalar() if row.chunk_id else None,
                "score": round(float(row.score), 6),
                "model_name": row.model_name,
                "model_version": row.model_version,
                "evidence": row.evidence_json,
                "graph_path": row.graph_path_json,
            } for row in rows],
        })
    payload = {
        "plan_id": plan.id,
        "project_version_id": version.id,
        "variant_type": plan.variant_type,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "course_ids": course_ids,
        "verification": (plan.metrics_json or {}).get("verification", {}),
        "evidence": evidence,
    }
    canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
    payload["sha256"] = hashlib.sha256(canonical).hexdigest()
    return payload
