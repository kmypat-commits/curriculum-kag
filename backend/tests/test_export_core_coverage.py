from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.api.export_api import append_core_coverage_sheet


def test_xlsx_shows_confirmed_and_unconfirmed_profile_blocks_as_snapshot():
    book = Workbook()
    book.active.title = "Curriculum Plan"
    metrics = {"core_coverage": {
        "enabled": True, "unique_core_credits": 5,
        "required_courses": {"requested": [17], "included": [17], "missing": []},
        "core_coverage": [
            {"block_id": "wood", "title": "Обработка древесины", "requirement": "required", "status": "covered", "supported_credits": 5, "selected_course_ids": [17]},
            {"block_id": "design", "title": "Проектирование мебели", "requirement": "preferred", "status": "unconfirmed", "supported_credits": 0, "unconfirmed_course_ids": [42]},
        ],
    }}

    append_core_coverage_sheet(book, metrics)

    sheet = book["Professional Core"]
    assert sheet["B2"].value == 5
    assert sheet["B3"].value == "1/1"
    assert "ручной правки" in sheet["A4"].value
    assert [cell.value for cell in sheet[7]] == [
        "Обработка древесины", "Обязательный", "Покрыт", 5, "17", "",
    ]
    assert [cell.value for cell in sheet[8]] == [
        "Проектирование мебели", "Предпочтительный", "Не подтверждён", 0, "", "42",
    ]


def test_legacy_xlsx_does_not_claim_profile_coverage():
    book = Workbook()
    append_core_coverage_sheet(book, {})
    assert "Professional Core" not in book.sheetnames


def test_block_title_is_exported_as_text_not_a_spreadsheet_formula():
    book = Workbook()
    append_core_coverage_sheet(book, {"core_coverage": {
        "enabled": True,
        "core_coverage": [{"title": '=HYPERLINK("https://example.invalid","open")', "status": "gap"}],
    }})
    cell = book["Professional Core"]["A7"]
    assert cell.value.startswith("=HYPERLINK")
    assert cell.data_type == "s"
    output = BytesIO()
    book.save(output)
    output.seek(0)
    saved = load_workbook(output)
    assert saved["Professional Core"]["A7"].data_type == "s"
