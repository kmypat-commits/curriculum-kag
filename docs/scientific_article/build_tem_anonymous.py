from pathlib import Path
from docx import Document

HERE = Path(__file__).parent
src = HERE / 'Curriculum_KAG_TEM_manuscript_official_template.docx'
out = HERE / 'Curriculum_KAG_TEM_manuscript_anonymous.docx'
doc = Document(src)

# The first four paragraphs are title, author, affiliation and correspondence
# in the official-template-derived manuscript.
replacements = {
    1: 'Anonymous Author(s)',
    2: 'Affiliations omitted for double-blind review',
    3: 'Corresponding-author details omitted for double-blind review',
}
for index, value in replacements.items():
    if index < len(doc.paragraphs):
        p = doc.paragraphs[index]
        for run in p.runs:
            run.text = ''
        if p.runs:
            p.runs[0].text = value
        else:
            p.add_run(value)

doc.core_properties.author = 'Anonymous'
doc.core_properties.last_modified_by = 'Anonymous'
doc.core_properties.comments = 'Double-blind review copy; author metadata omitted.'
doc.save(out)
print(out)
