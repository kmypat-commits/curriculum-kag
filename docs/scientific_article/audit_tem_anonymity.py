"""Check that the TEM review copy does not expose author identity."""
from pathlib import Path
import json
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parent
path = ROOT / "Curriculum_KAG_TEM_manuscript_anonymous.docx"
doc = Document(path)
text = "\n".join(p.text for p in doc.paragraphs).lower()
forbidden = ("murat kozhanov", "kozhanov", "@ktu.edu.kz", "corresponding author")
found = [token for token in forbidden if token in text]
core = doc.core_properties
metadata = " ".join(
    str(value or "") for value in (core.author, core.last_modified_by, core.comments)
).lower()
found.extend(token for token in forbidden if token in metadata)

with ZipFile(path) as archive:
    names = set(archive.namelist())
    hidden_review_artifacts = sorted(
        name for name in names
        if "comments" in name.lower() or "track" in name.lower() or "people" in name.lower()
    )

result = {
    "path": str(path),
    "forbidden_identity_tokens": sorted(set(found)),
    "hidden_review_artifacts": hidden_review_artifacts,
    "first_paragraphs": [p.text for p in doc.paragraphs[:4]],
    "pass": not found and not hidden_review_artifacts,
}
(ROOT / "tem_anonymity_audit.json").write_text(
    json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(result)
if not result["pass"]:
    raise SystemExit(1)
