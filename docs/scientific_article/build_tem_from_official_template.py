from copy import deepcopy
from pathlib import Path
import re
from docx import Document
from docx.oxml.ns import qn

HERE = Path(__file__).parent
TEMPLATE = HERE / 'TEM_official_template.docx'
SOURCE = HERE / 'Curriculum_KAG_TEM_manuscript_draft.docx'
OUT = HERE / 'Curriculum_KAG_TEM_manuscript_official_template.docx'

template = Document(TEMPLATE)
source = Document(SOURCE)

# Retain the official document body section properties and the internal section
# break. Remove only sample content; styles, numbering, headers, footers and
# relationships remain inherited from the official template.
body = template.element.body
section_break = None
for child in list(body):
    if child.tag == qn('w:p') and child.find('.//' + qn('w:sectPr')) is not None:
        section_break = deepcopy(child)
    if child.tag != qn('w:sectPr'):
        body.remove(child)

def style_for(text, source_style):
    if text.startswith('Curriculum-KAG:'):
        return 'ICEST_Title'
    if text.startswith('Abstract') or text.startswith('Keywords'):
        return 'ICEST_Abstract'
    if re.match(r'^\d+\.', text) or re.match(r'^\d+\.\d+', text):
        return 'ICEST_Title'
    if text in {'Data availability', 'Code availability', 'Ethics and conflicts',
                'Generative AI disclosure', 'References'}:
        return 'ICEST_Title'
    if text.startswith('[') and re.match(r'^\[\d+\]', text):
        return 'List Paragraph'
    if source_style.startswith('Heading'):
        return 'ICEST_Title'
    if source_style == 'List Bullet':
        return 'List Paragraph'
    return 'ICEST_Normal'

seen_keywords = False
for child in source.element.body:
    if child.tag == qn('w:sectPr'):
        continue
    if child.tag == qn('w:p'):
        source_para = next((p for p in source.paragraphs if p._p is child), None)
        # lxml itertext() can expose duplicated text from compatibility markup
        # in this generated DOCX; python-docx's paragraph text is canonical.
        text = (source_para.text if source_para is not None else '').strip()
        if not text:
            continue
        # Map the source paragraph to a real template style and retain only
        # semantic text; the template supplies the typography and geometry.
        source_style = source_para.style.name if source_para is not None else 'Normal'
        p = template.add_paragraph(style=style_for(text, source_style))
        p.add_run(text)
        if text.startswith('Keywords'):
            seen_keywords = True
            if section_break is not None:
                p._p.addnext(deepcopy(section_break))
    elif child.tag == qn('w:tbl'):
        body.insert(body.index(body[-1]), deepcopy(child))

template.core_properties.title = 'Curriculum-KAG: Expert-Grounded and Constraint-Aware Curriculum Synthesis from Learning Outcomes'
template.core_properties.author = 'Murat Kozhanov'
template.core_properties.subject = 'TEM Journal manuscript'
template.core_properties.keywords = 'curriculum design, EPVO, SBERT, learning outcomes, expert evidence'
template.save(OUT)
print(OUT)
