"""Deterministic pre-submission audit for the TEM Journal DOCX package."""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent
FILES = {
    "template": ROOT / "TEM_official_template.docx",
    "full_manuscript": ROOT / "Curriculum_KAG_TEM_manuscript_official_template.docx",
    "anonymous_manuscript": ROOT / "Curriculum_KAG_TEM_manuscript_anonymous.docx",
    "supplement": ROOT / "TEM_SUPPLEMENTARY_MATERIALS_EN.md",
    "checklist": ROOT / "TEM_SUBMISSION_CHECKLIST_EN.md",
    "frozen_table": ROOT / "TEM_FROZEN_RESULTS_TABLE_EN.md",
    "checksum_manifest": ROOT / "TEM_CHECKSUM_MANIFEST.json",
    "blinded_protocol": ROOT / "TEM_BLINDED_REVIEW_PROTOCOL_EN.md",
    "anonymity_audit": ROOT / "tem_anonymity_audit.json",
    "upload_instructions": ROOT / "TEM_UPLOAD_INSTRUCTIONS_RU.md",
    "text_audit": ROOT / "tem_text_audit.json",
    "citation_audit": ROOT / "tem_citation_audit.json",
    "readiness_report": ROOT / "TEM_READINESS_REPORT_RU.md",
}


def words(text: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", text, flags=re.UNICODE))


def inspect_docx(path: Path) -> dict:
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    abstract = ""
    if "Abstract —" in text and "Keywords —" in text:
        abstract = text.split("Abstract —", 1)[1].split("Keywords —", 1)[0].strip()
    keyword_block = ""
    if "Keywords —" in text and "1. Introduction" in text:
        keyword_block = text.split("Keywords —", 1)[1].split("1. Introduction", 1)[0].strip()
    placeholders = [
        p.text for p in doc.paragraphs
        if any(token in p.text for token in ("[Insert", "TODO", "TBD", "Lorem ipsum"))
    ]
    return {
        "exists": path.exists(),
        "bytes": path.stat().st_size if path.exists() else 0,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "words": words(text),
        "abstract_words": words(abstract),
        "keywords": keyword_block,
        "placeholder_paragraphs": placeholders,
        "has_title": bool(doc.paragraphs and doc.paragraphs[0].text.strip()),
    }


def main() -> None:
    result = {
        "package": {name: str(path) for name, path in FILES.items()},
        "documents": {
            name: inspect_docx(path)
            for name, path in FILES.items()
            if path.suffix.lower() == ".docx"
        },
        "supporting_files": {
            name: {"exists": path.exists(), "bytes": path.stat().st_size if path.exists() else 0}
            for name, path in FILES.items()
            if path.suffix.lower() != ".docx"
        },
    }
    result["pass"] = all(
        item["exists"] and not item["placeholder_paragraphs"]
        for item in result["documents"].values()
    ) and all(item["exists"] for item in result["supporting_files"].values())
    out = ROOT / "tem_package_audit.json"
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
